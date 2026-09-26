# -*- coding: utf-8 -*-
"""Approximate PPTX renderer for layout QA when Office/LibreOffice is unavailable."""

from io import BytesIO
from pathlib import Path
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN


ROOT = Path(__file__).resolve().parent.parent
enhanced = list((ROOT / "reports").glob("*11页数据增强版.pptx"))
concise = list((ROOT / "reports").glob("*10页精简版.pptx"))
academic = list((ROOT / "reports").glob("*学术重构版.pptx"))
PPTX = Path(sys.argv[1]) if len(sys.argv) > 1 else (
    enhanced[0] if enhanced else (
        concise[0] if concise else (
            academic[0] if academic else next((ROOT / "reports").glob("*开题答辩*图检索*.pptx"))
        )
    )
)
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "reports" / "proposal_preview"
OUT.mkdir(exist_ok=True)

W, H = 1200, 900
prs = Presentation(str(PPTX))
sx, sy = W / prs.slide_width, H / prs.slide_height
FONT_REG = Path("C:/Windows/Fonts/msyh.ttc")
FONT_BOLD = Path("C:/Windows/Fonts/msyhbd.ttc")


def rgb(color, default=(255, 255, 255)):
    try:
        value = color.rgb
        if value is None:
            return default
        return tuple(value)
    except Exception:
        return default


def shape_fill(shape, default=(255, 255, 255)):
    try:
        if shape.fill.type is None:
            return None
        return rgb(shape.fill.fore_color, default)
    except Exception:
        return default


def shape_line(shape, default=(210, 215, 220)):
    try:
        if shape.line.fill.type is None:
            return None
        return rgb(shape.line.color, default)
    except Exception:
        return default


def font(size, bold=False):
    return ImageFont.truetype(str(FONT_BOLD if bold and FONT_BOLD.exists() else FONT_REG),
                              max(8, int(size * 96 / 72)))


def wrap(draw, text, ft, width):
    lines = []
    for raw in text.split("\n"):
        if not raw:
            lines.append("")
            continue
        current = ""
        for ch in raw:
            candidate = current + ch
            if current and draw.textlength(candidate, font=ft) > width:
                lines.append(current)
                current = ch
            else:
                current = candidate
        lines.append(current)
    return lines


def draw_text_frame(draw, shape, x, y, w, h):
    tf = shape.text_frame
    ml = int(tf.margin_left * sx); mr = int(tf.margin_right * sx)
    mt = int(tf.margin_top * sy); mb = int(tf.margin_bottom * sy)
    blocks = []
    for para in tf.paragraphs:
        text = para.text
        if not text:
            continue
        run = next((r for r in para.runs if r.text), None)
        size = run.font.size.pt if run is not None and run.font.size else 12
        bold = bool(run.font.bold) if run is not None else False
        color = rgb(run.font.color, (32, 41, 51)) if run is not None else (32, 41, 51)
        ft = font(size, bold)
        lines = wrap(draw, text, ft, max(10, w - ml - mr))
        line_h = max(10, int(size * 96 / 72 * 1.18))
        blocks.append((para.alignment, lines, ft, color, line_h))
    total_h = sum(len(lines) * lh + 2 for _, lines, _, _, lh in blocks)
    anchor = tf.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        cy = y + max(mt, (h - total_h) // 2)
    elif anchor == MSO_ANCHOR.BOTTOM:
        cy = y + h - total_h - mb
    else:
        cy = y + mt
    for align, lines, ft, color, line_h in blocks:
        for line in lines:
            tw = draw.textlength(line, font=ft)
            if align == PP_ALIGN.CENTER:
                cx = x + (w - tw) / 2
            elif align == PP_ALIGN.RIGHT:
                cx = x + w - mr - tw
            else:
                cx = x + ml
            draw.text((cx, cy), line, font=ft, fill=color)
            cy += line_h
        cy += 2


def draw_shape(draw, image, shape):
    x, y = int(shape.left * sx), int(shape.top * sy)
    w, h = max(1, int(shape.width * sx)), max(1, int(shape.height * sy))
    bbox = (x, y, x + w, y + h)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            pic = Image.open(BytesIO(shape.image.blob)).convert("RGBA")
            pic.thumbnail((w, h), Image.Resampling.LANCZOS)
            image.alpha_composite(pic, (x, y))
        except Exception:
            pass
        return
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        draw.line((x, y, x + w, y + h), fill=shape_line(shape, (170, 180, 190)) or (170, 180, 190), width=2)
        return
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        table = shape.table
        cy = y
        for row in table.rows:
            rh = int(row.height * sy); cx = x
            for j, cell in enumerate(row.cells):
                cw = int(table.columns[j].width * sx)
                fill = shape_fill(cell, (255, 255, 255))
                draw.rectangle((cx, cy, cx + cw, cy + rh), fill=fill, outline=(210, 217, 224), width=1)
                proxy = type("CellProxy", (), {"text_frame": cell.text_frame})()
                draw_text_frame(draw, proxy, cx, cy, cw, rh)
                cx += cw
            cy += rh
        return
    fill = shape_fill(shape, (255, 255, 255))
    outline = shape_line(shape, (210, 217, 224))
    try:
        name = str(shape.auto_shape_type)
    except Exception:
        name = ""
    if "OVAL" in name:
        draw.ellipse(bbox, fill=fill, outline=outline, width=1)
    elif "ARROW" in name:
        if w >= h:
            pts = [(x, y + h * 0.25), (x + w * 0.60, y + h * 0.25),
                   (x + w * 0.60, y), (x + w, y + h / 2),
                   (x + w * 0.60, y + h), (x + w * 0.60, y + h * 0.75),
                   (x, y + h * 0.75)]
        else:
            pts = [(x + w * 0.25, y), (x + w * 0.75, y),
                   (x + w * 0.75, y + h * 0.60), (x + w, y + h * 0.60),
                   (x + w / 2, y + h), (x, y + h * 0.60),
                   (x + w * 0.25, y + h * 0.60)]
        draw.polygon(pts, fill=fill)
    elif "ROUNDED_RECTANGLE" in name:
        draw.rounded_rectangle(bbox, radius=max(3, min(w, h) // 10), fill=fill, outline=outline, width=1)
    else:
        draw.rectangle(bbox, fill=fill, outline=outline, width=1)
    if getattr(shape, "has_text_frame", False) and shape.text.strip():
        draw_text_frame(draw, shape, x, y, w, h)


thumbs = []
for idx, slide in enumerate(prs.slides, 1):
    canvas = Image.new("RGBA", (W, H), (248, 250, 252, 255))
    draw = ImageDraw.Draw(canvas)
    for shape in slide.shapes:
        draw_shape(draw, canvas, shape)
    path = OUT / f"slide_{idx:02d}.png"
    canvas.convert("RGB").save(path, quality=95)
    thumb = canvas.convert("RGB").resize((400, 300), Image.Resampling.LANCZOS)
    thumbs.append(thumb)

sheet = Image.new("RGB", (1600, 1200), "white")
for i, thumb in enumerate(thumbs):
    sheet.paste(thumb, ((i % 4) * 400, (i // 4) * 300))
sheet_path = OUT / "contact_sheet.png"
sheet.save(sheet_path, quality=95)
print(sheet_path)
