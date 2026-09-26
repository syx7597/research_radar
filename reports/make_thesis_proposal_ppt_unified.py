# -*- coding: utf-8 -*-
"""Generate a concise 10-slide academic thesis proposal deck."""

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "PPT模板1.pptx"
FIG = ROOT / "reports" / "academic_figures"
OUTPUT = ROOT / "reports" / "硕士学位论文开题答辩_图检索策略学习_统一方法精简版.pptx"

NAVY = RGBColor(0x18, 0x36, 0x57)
BLUE = RGBColor(0x2F, 0x64, 0x91)
TEAL = RGBColor(0x4F, 0x7D, 0x7B)
GOLD = RGBColor(0xB4, 0x88, 0x38)
INK = RGBColor(0x24, 0x2D, 0x36)
GRAY = RGBColor(0x68, 0x73, 0x7D)
MID = RGBColor(0xB5, 0xC1, 0xCC)
LINE = RGBColor(0xD8, 0xE0, 0xE7)
PALE_BLUE = RGBColor(0xEC, 0xF1, 0xF5)
PALE_TEAL = RGBColor(0xEB, 0xF2, 0xF1)
PALE_GOLD = RGBColor(0xF6, 0xF1, 0xE6)
BG = RGBColor(0xFA, 0xFB, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"
MONO = "Consolas"


prs = Presentation(str(TEMPLATE))
logo_shape = next(sh for sh in prs.slides[0].shapes if sh.shape_type == 13)
LOGO = BytesIO(logo_shape.image.blob)
while len(prs.slides):
    sid = prs.slides._sldIdLst[0]
    prs.part.drop_rel(sid.rId)
    del prs.slides._sldIdLst[0]
BLANK = prs.slide_layouts[6]
prs.core_properties.title = "面向复杂雷达情报问答的图检索策略学习方法研究"
prs.core_properties.subject = "硕士学位论文开题答辩"


def set_typeface(run, name=FONT):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", name)


def style_run(run, size=12, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_typeface(run, name)


def text(slide, value, x, y, w, h, size=12, bold=False, color=INK,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03, name=FONT):
    obj = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = obj.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    style_run(r, size, bold, color, name)
    return obj


def rect(slide, x, y, w, h, fill=WHITE, border=LINE,
         kind=MSO_SHAPE.RECTANGLE, width=0.8):
    obj = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = fill
    if border is None:
        obj.line.fill.background()
    else:
        obj.line.color.rgb = border
        obj.line.width = Pt(width)
    obj.shadow.inherit = False
    return obj


def box(slide, value, x, y, w, h, fill=WHITE, border=LINE, size=11,
        bold=False, color=INK, align=PP_ALIGN.CENTER, width=0.8):
    obj = rect(slide, x, y, w, h, fill, border, width=width)
    tf = obj.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    style_run(r, size, bold, color)
    return obj


def connector(slide, x1, y1, x2, y2, color=MID, width=1.0):
    obj = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    obj.line.color.rgb = color
    obj.line.width = Pt(width)
    return obj


def arrow(slide, x, y, w, h, color=BLUE):
    return rect(slide, x, y, w, h, color, None, MSO_SHAPE.RIGHT_ARROW)


def add_picture(slide, path, x, y, w=None, h=None):
    kwargs = {}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)


def footer(slide, page):
    connector(slide, 0.45, 7.13, 9.56, 7.13, LINE, 0.6)
    text(slide, "硕士学位论文开题答辩  |  图检索策略学习", 0.47, 7.18, 5.8, 0.17,
         7.8, color=GRAY)
    text(slide, f"{page:02d}", 9.08, 7.17, 0.38, 0.18, 8.0, True, GRAY,
         PP_ALIGN.RIGHT)


def new_slide(title_value, section, page):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    rect(slide, 0, 0, 0.07, 7.5, NAVY, None)
    text(slide, section, 0.47, 0.18, 4.0, 0.22, 8.8, True, TEAL)
    text(slide, title_value, 0.47, 0.48, 8.2, 0.50, 20, True, NAVY)
    connector(slide, 0.47, 1.08, 9.55, 1.08, MID, 0.8)
    LOGO.seek(0)
    slide.shapes.add_picture(LOGO, Inches(7.75), Inches(0.12), width=Inches(2.03))
    footer(slide, page)
    return slide


def conclusion(slide, value, y=6.48, fill=PALE_TEAL, border=TEAL):
    rect(slide, 0.62, y, 8.76, 0.48, fill, border, width=0.8)
    text(slide, "结论", 0.80, y + 0.11, 0.48, 0.20, 8.8, True, border,
         PP_ALIGN.CENTER)
    connector(slide, 1.38, y + 0.10, 1.38, y + 0.38, border, 0.9)
    text(slide, value, 1.55, y + 0.09, 7.56, 0.24, 10.2, True, NAVY,
         PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


def table(slide, headers, rows, x, y, w, h, widths, size=8.5,
          highlight_rows=()):
    tbl = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    for i, value in enumerate(widths):
        tbl.columns[i].width = Inches(value)
    for j, value in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.04)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = value
        style_run(r, size, True, WHITE)
    for i, row in enumerate(rows, 1):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            if i - 1 in highlight_rows:
                cell.fill.fore_color.rgb = PALE_GOLD
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else PALE_BLUE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.04)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j in (0, len(row) - 1) else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(value)
            style_run(r, size, i - 1 in highlight_rows and j == 0, INK)
    return tbl


def small_label(slide, value, x, y, w, color=TEAL):
    rect(slide, x, y + 0.03, 0.04, 0.22, color, None)
    text(slide, value, x + 0.12, y, w - 0.12, 0.27, 9.4, True, color)


# 1 Cover ------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG
rect(s, 0, 0, 0.12, 7.5, NAVY, None)
rect(s, 0.65, 1.18, 0.80, 0.05, TEAL, None)
text(s, "硕士学位论文开题答辩", 0.66, 0.70, 3.6, 0.34, 12, True, TEAL)
text(s, "面向复杂雷达情报问答的\n图检索策略学习方法研究",
     0.66, 1.48, 8.45, 1.56, 27, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
text(s, "Graph Retrieval Policy Learning for Complex Radar Intelligence Question Answering",
     0.70, 3.20, 8.20, 0.40, 10.5, False, GRAY)
connector(s, 0.70, 4.18, 8.90, 4.18, LINE, 1.0)
text(s, "研究主线", 0.70, 4.48, 0.88, 0.24, 9.5, True, TEAL)
text(s, "复杂组合检索  →  可执行策略空间  →  执行反馈优化  →  组合泛化验证",
     1.72, 4.43, 7.15, 0.34, 12, True, NAVY)
text(s, "汇报人：__________    专业：__________    指导教师：__________",
     0.70, 5.65, 6.90, 0.32, 10.5, False, INK)
text(s, "2026 年 8 月", 0.70, 6.15, 2.0, 0.30, 10, False, GRAY)
LOGO.seek(0)
s.shapes.add_picture(LOGO, Inches(7.58), Inches(5.78), width=Inches(2.14))
text(s, "01", 9.12, 7.18, 0.35, 0.18, 8.0, True, GRAY, PP_ALIGN.RIGHT)


# 2 Problem ----------------------------------------------------------------
s = new_slide("复杂雷达问答本质上是组合检索与精确执行问题", "研究背景与问题定义", 2)
box(s, "示例任务（结构示意）\n检索满足多重条件的雷达，关联其平台与研制方，完成集合筛选、计数与比较",
    0.72, 1.36, 8.56, 0.74, PALE_BLUE, BLUE, 12.0, True, NAVY)
small_label(s, "异构知识环境", 0.74, 2.40, 1.50)
for i, (a, b) in enumerate([
    ("知识图谱", "实体 · 关系 · 属性"),
    ("叙述文本", "原理 · 结构 · 背景"),
    ("工具接口", "路径 · 集合 · 聚合"),
]):
    y = 2.78 + i * 0.74
    box(s, a, 0.80, y, 1.24, 0.50, WHITE, BLUE, 10.4, True, NAVY)
    text(s, b, 2.18, y + 0.13, 1.45, 0.23, 9.0, False, GRAY)
connector(s, 3.55, 2.72, 3.55, 4.97, MID, 0.9)
small_label(s, "组合查询链", 4.05, 2.40, 1.50)
steps = ["实体对齐", "约束筛选", "关系扩展", "交并差", "计数 / 比较"]
for i, v in enumerate(steps):
    x = 4.04 + (i % 3) * 1.62
    y = 2.82 + (i // 3) * 1.08
    box(s, v, x, y, 1.25, 0.48, WHITE, BLUE if i < 3 else TEAL,
        9.8, True, NAVY)
    if i in (0, 1, 3):
        arrow(s, x + 1.31, y + 0.16, 0.22, 0.16, BLUE)
text(s, "中间结果与变量依赖决定后续动作", 4.08, 5.08, 4.62, 0.28,
     10.3, True, TEAL, PP_ALIGN.CENTER)
conclusion(s, "问题难点不在一次召回，而在多步检索、精确计算与状态驱动决策。")


# 3 Gap --------------------------------------------------------------------
s = new_slide("现有方法的核心缺口是计划表达与长程学习", "研究现状与科学问题", 3)
table(s, ["方法范式", "主要能力", "复杂组合问答中的关键缺口"], [
    ["混合 RAG", "关键词与语义证据召回", "难以稳定表达集合、计数及多步状态"],
    ["图路径推理", "关系路径规划与子图检索", "集合聚合、参数约束与错误修正能力有限"],
    ["Prompt / ReAct Agent", "灵活调用工具并迭代观察", "自由动作空间下参数生成与长程决策不稳定"],
    ["本文：可执行策略学习", "类型化计划 + 确定性执行 + 反馈优化", "面向计划表达与信用分配建立统一方法"],
], 0.66, 1.40, 8.68, 3.20, [1.65, 2.62, 4.41], 9.2, highlight_rows=(3,))
small_label(s, "两个递进研究问题", 0.72, 4.94, 2.10)
box(s, "策略空间如何设计？\n将任务分解、算子、参数和变量依赖组织为可执行计划",
    0.78, 5.28, 3.92, 0.82, WHITE, BLUE, 10.1, True, NAVY)
box(s, "策略如何有效学习？\n利用执行过程反馈缓解终局奖励稀疏与长程信用分配",
    5.00, 5.28, 4.22, 0.82, WHITE, TEAL, 10.1, True, NAVY)
conclusion(s, "研究目标是学习多步可执行图检索策略，而非训练一次工具路由器。", 6.43)


# 4 Unified framework -------------------------------------------------------
s = new_slide("统一方法：可执行图检索策略学习", "总体研究思路与技术路线", 4)
add_picture(s, FIG / "unified_framework.png", 0.55, 1.32, w=8.92)
conclusion(s, "策略空间定义可表达的决策；执行反馈用于优化同一策略的规划、修正与终止。", 6.47)


# 5 Policy space ------------------------------------------------------------
s = new_slide("策略空间构造：将复杂查询表示为类型安全的可执行计划", "统一方法  |  策略表示与执行", 5)
add_picture(s, FIG / "policy_space.png", 0.57, 1.34, w=8.86)
conclusion(s, "LLM 负责语义决策，执行器负责确定性计算，实现规划能力与计算可靠性的职责解耦。", 6.47)


# 6 Policy learning ---------------------------------------------------------
s = new_slide("策略优化：利用执行反馈学习多步决策与修正", "统一方法  |  策略训练与优化", 6)
add_picture(s, FIG / "policy_learning.png", 0.58, 1.34, w=8.84)
conclusion(s, "优化对象为已定义的类型化检索策略 πθ，训练信号来自其完整执行轨迹。", 6.47)


# 7 Foundation --------------------------------------------------------------
s = new_slide("现有知识环境与轨迹数据已支撑策略学习实验", "已完成研究基础", 7)
table(s, ["层次", "已完成资产", "规模 / 结果", "对论文研究的作用"], [
    ["知识环境", "RadarKG + 叙述文本库", "10,744 实体；22,241 边；1,432 段", "提供可执行图文知识环境"],
    ["复杂问答", "Train / Dev / Test", "4,277 / 582 / 542", "支撑训练、选择与独立测试"],
    ["执行轨迹", "成功轨迹 / Oracle 回放", "4,277 / 5,401", "用于 SFT 冷启动与轨迹评测"],
    ["执行器", "10 类类型化算子", "72 题原型：计划 100%；执行 100%；答案 98.6%", "验证工具协议与执行链路"],
    ["知识质量", "分层人工抽样 n=173", "加权准确率 85.1% [75.1, 95.0]", "界定知识环境的可验证质量"],
], 0.58, 1.40, 8.84, 4.78, [1.15, 2.05, 2.72, 2.92], 8.2)
text(s, "状态说明：以上均为已完成数据与原型结果；SFT/RL 的正式 2×2 对照为后续拟开展实验。",
     0.78, 6.24, 8.44, 0.20, 8.4, False, GRAY, PP_ALIGN.CENTER)
conclusion(s, "知识环境、复杂问答、成功轨迹和确定性执行器共同构成策略学习的实验基础。", 6.52)


# 8 Prior evidence ----------------------------------------------------------
s = new_slide("前期实验表明：结构化计划有效，单步路由并非主要瓶颈", "已有实验与研究动机", 8)
add_picture(s, FIG / "representative_results.png", 0.60, 1.42, w=4.76)
table(s, ["诊断实验", "结果", "研究启示"], [
    ["LLM / Oracle 路由", "88.6 / 88.2\nΔ=-0.4 pp", "一次路由已接近上限，关键在计划执行"],
    ["Qwen Planner 微调", "31.5 → 61.7", "成功轨迹能够提升规划能力"],
    ["类型化策略 vs Qwen-FT", "+26.9 pp\n[22.6, 31.3]", "结构化计划表达仍有显著价值"],
    ["KQA Pro 初测 n=502", "Δ=-0.2 pp\n[-2.4, 2.0]", "固定策略存在任务依赖，需学习与环境适配"],
], 5.52, 1.43, 3.86, 4.45, [1.15, 0.86, 1.85], 7.6)
text(s, "RadarKG-QA-499，点为总体准确率，线为 95% CI",
     0.75, 5.70, 4.40, 0.20, 7.8, False, GRAY, PP_ALIGN.CENTER)
conclusion(s, "前期证据将研究重点定位到多步计划表示与执行反馈学习。", 6.43)


# 9 Evaluation --------------------------------------------------------------
s = new_slide("2×2 对照独立验证策略表示与策略优化", "实验设计与公开验证", 9)
small_label(s, "核心因果对照", 0.63, 1.34, 1.60)
text(s, "扁平动作空间", 2.10, 1.72, 1.55, 0.24, 9.5, True, GRAY, PP_ALIGN.CENTER)
text(s, "类型化策略空间", 4.20, 1.72, 1.72, 0.24, 9.5, True, BLUE, PP_ALIGN.CENTER)
text(s, "SFT", 0.78, 2.34, 0.55, 0.24, 9.5, True, GRAY, PP_ALIGN.CENTER)
text(s, "RL", 0.78, 3.40, 0.55, 0.24, 9.5, True, GRAY, PP_ALIGN.CENTER)
box(s, "扁平 SFT\n成功轨迹模仿", 1.70, 2.12, 2.00, 0.76, PALE_BLUE, MID, 10.0, True, NAVY)
box(s, "类型化 SFT\n结构化动作 + 轨迹", 4.02, 2.12, 2.00, 0.76, WHITE, BLUE, 10.0, True, NAVY)
box(s, "扁平 RL\n终局 / 执行奖励", 1.70, 3.18, 2.00, 0.76, PALE_BLUE, MID, 10.0, True, NAVY)
box(s, "完整方法 RL\n类型化策略 + 多粒度反馈", 4.02, 3.18, 2.00, 0.76, PALE_GOLD, GOLD, 10.0, True, NAVY)
arrow(s, 3.72, 2.39, 0.24, 0.17, BLUE)
text(s, "验证策略空间构造", 2.58, 2.92, 2.72, 0.20, 8.0, True, BLUE, PP_ALIGN.CENTER)
connector(s, 5.02, 2.93, 5.02, 3.16, GOLD, 1.5)
text(s, "验证执行反馈优化", 5.18, 2.95, 1.22, 0.38, 8.0, True, GOLD)

small_label(s, "验证层与公开数据", 6.48, 1.34, 2.20)
table(s, ["数据集", "主要作用"], [
    ["RadarKG-QA", "IID / 实体隔离 / 组合隔离 / 错误恢复"],
    ["KQA Pro", "程序化组合问答与可执行计划验证"],
    ["MetaQA", "1 / 2 / 3-hop 跨领域多跳验证"],
], 6.48, 1.72, 2.90, 2.22, [1.02, 1.88], 7.7)
text(s, "组合隔离：训练与测试使用不同算子组合，检验策略学习而非模板记忆。",
     6.55, 4.12, 2.70, 0.70, 8.4, True, NAVY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
rect(s, 0.72, 5.15, 8.55, 0.72, PALE_BLUE, BLUE)
text(s, "分层报告", 0.94, 5.35, 0.72, 0.22, 9.0, True, BLUE)
text(s, "环境 Oracle  →  策略执行成功率  →  最终答案 EM / F1  →  步数、重复与恢复成本",
     1.78, 5.30, 7.20, 0.30, 10.2, True, NAVY, PP_ALIGN.CENTER)
conclusion(s, "横向对照验证动作空间设计，纵向对照验证执行反馈学习，组合隔离检验泛化。", 6.43)


# 10 Contributions and plan -------------------------------------------------
s = new_slide("论文围绕一个统一方法形成两项递进贡献", "创新点、论文结构与研究计划", 10)
table(s, ["递进环节", "核心贡献", "解决问题", "主要验证"], [
    ["策略空间构造", "类型化可执行图检索决策建模", "动作空间过大、参数与变量依赖不稳定", "扁平 SFT → 类型化 SFT"],
    ["执行反馈优化", "基于多粒度执行反馈的检索策略学习", "终局奖励稀疏与长程信用分配", "类型化 SFT → 完整方法 RL"],
], 0.62, 1.38, 8.76, 1.72, [1.35, 2.55, 2.73, 2.13], 8.5, highlight_rows=(0, 1))

small_label(s, "论文章节", 0.66, 3.44, 1.30)
chapters = [
    ("第 3 章", "知识环境与任务"),
    ("第 4 章", "策略空间构造"),
    ("第 5 章", "执行反馈优化"),
    ("第 6 章", "实验与泛化验证"),
]
for i, (a, b) in enumerate(chapters):
    x = 0.72 + i * 2.18
    box(s, f"{a}\n{b}", x, 3.82, 1.74, 0.66, WHITE, BLUE if i < 3 else TEAL,
        9.5, True, NAVY)
    if i < 3:
        arrow(s, x + 1.82, 4.07, 0.22, 0.16, BLUE)

small_label(s, "研究计划", 0.66, 4.82, 1.30)
phases = [
    ("1–2 月", "任务 / 数据冻结"),
    ("3–4 月", "环境与基线"),
    ("5–7 月", "SFT / RL / 消融"),
    ("8–10 月", "公开验证 / 写作"),
]
for i, (a, b) in enumerate(phases):
    x = 0.72 + i * 2.18
    rect(s, x, 5.18, 1.74, 0.62, PALE_BLUE if i != 2 else PALE_GOLD,
         BLUE if i != 2 else GOLD)
    text(s, a, x + 0.08, 5.28, 0.52, 0.22, 8.3, True, TEAL, PP_ALIGN.CENTER)
    text(s, b, x + 0.62, 5.24, 1.02, 0.28, 8.2, True, NAVY,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
conclusion(s, "形成可执行、可学习、可独立验证且具组合泛化能力的图检索策略方法。", 6.43, PALE_GOLD, GOLD)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUTPUT))
print(f"saved: {OUTPUT}")
print(f"slides: {len(prs.slides)}; size={prs.slide_width / 914400:.2f}x{prs.slide_height / 914400:.2f}")
