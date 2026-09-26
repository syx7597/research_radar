# -*- coding: utf-8 -*-
"""Build the final 10-slide proposal deck plus four backup slides."""

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "reports" / "硕士学位论文开题答辩_图检索策略学习_统一方法精简版.pptx"
OUTPUT = ROOT / "reports" / "硕士学位论文开题答辩_图检索策略学习_最终版.pptx"
RESULT_FIG = ROOT / "reports" / "academic_figures" / "representative_results.png"

NAVY = RGBColor(0x18, 0x36, 0x57)
BLUE = RGBColor(0x2F, 0x64, 0x91)
TEAL = RGBColor(0x4F, 0x7D, 0x7B)
GOLD = RGBColor(0xB4, 0x88, 0x38)
INK = RGBColor(0x24, 0x2D, 0x36)
GRAY = RGBColor(0x68, 0x73, 0x7D)
MID = RGBColor(0xB5, 0xC1, 0xCC)
LINE = RGBColor(0xD8, 0xE0, 0xE7)
PALE_BLUE = RGBColor(0xEC, 0xF1, 0xF5)
PALE_TEAL = RGBColor(0xEB, 0xF2, 0xF1)
PALE_GOLD = RGBColor(0xF6, 0xF1, 0xE6)
BG = RGBColor(0xFA, 0xFB, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"
MONO = "Consolas"


prs = Presentation(str(BASE))
logo_shape = next(sh for sh in prs.slides[0].shapes if sh.shape_type == 13)
LOGO = BytesIO(logo_shape.image.blob)
BLANK = prs.slide_layouts[6]
prs.core_properties.title = "面向复杂雷达情报问答的图检索策略学习方法研究"
prs.core_properties.subject = "硕士学位论文开题答辩最终版"


def set_typeface(run, name=FONT):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", name)


def style_run(run, size=12, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_typeface(run, name)


def text(slide, value, x, y, w, h, size=12, bold=False, color=INK,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03, name=FONT):
    obj = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = obj.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    style_run(r, size, bold, color, name)
    return obj


def rect(slide, x, y, w, h, fill=WHITE, border=LINE,
         kind=MSO_SHAPE.RECTANGLE, width=0.8):
    obj = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = fill
    if border is None:
        obj.line.fill.background()
    else:
        obj.line.color.rgb = border
        obj.line.width = Pt(width)
    obj.shadow.inherit = False
    return obj


def box(slide, value, x, y, w, h, fill=WHITE, border=LINE, size=10.5,
        bold=False, color=INK, align=PP_ALIGN.CENTER, width=0.8, name=FONT):
    obj = rect(slide, x, y, w, h, fill, border, width=width)
    tf = obj.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    style_run(r, size, bold, color, name)
    return obj


def line(slide, x1, y1, x2, y2, color=MID, width=1.0):
    obj = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    obj.line.color.rgb = color
    obj.line.width = Pt(width)
    return obj


def arrow(slide, x, y, w, h, color=BLUE, direction="right"):
    kinds = {
        "right": MSO_SHAPE.RIGHT_ARROW,
        "left": MSO_SHAPE.LEFT_ARROW,
        "down": MSO_SHAPE.DOWN_ARROW,
        "up": MSO_SHAPE.UP_ARROW,
    }
    return rect(slide, x, y, w, h, color, None, kinds[direction])


def clear_slide(slide):
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)


def footer(slide, page, backup=False):
    line(slide, 0.45, 7.13, 9.56, 7.13, LINE, 0.6)
    label = "硕士学位论文开题答辩  |  图检索策略学习"
    if backup:
        label += "  |  BACKUP"
    text(slide, label, 0.47, 7.18, 6.1, 0.17, 7.8, color=GRAY)
    page_text = f"A{page - 10}" if backup else f"{page:02d}"
    text(slide, page_text, 9.02, 7.17, 0.45, 0.18, 8.0, True, GRAY,
         PP_ALIGN.RIGHT)


def setup(slide, title_value, section, page, backup=False):
    clear_slide(slide)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    rect(slide, 0, 0, 0.07, 7.5, NAVY, None)
    text(slide, section, 0.47, 0.18, 4.2, 0.22, 8.8, True, TEAL)
    text(slide, title_value, 0.47, 0.48, 8.25, 0.50, 20, True, NAVY)
    line(slide, 0.47, 1.08, 9.55, 1.08, MID, 0.8)
    LOGO.seek(0)
    slide.shapes.add_picture(LOGO, Inches(7.75), Inches(0.12), width=Inches(2.03))
    footer(slide, page, backup)
    return slide


def new_backup(title_value, section, page):
    return setup(prs.slides.add_slide(BLANK), title_value, section, page, True)


def takeaway(slide, value, y=6.48, accent=TEAL):
    rect(slide, 0.64, y, 0.045, 0.34, accent, None)
    text(slide, value, 0.83, y + 0.02, 8.36, 0.28, 10.2, True, NAVY,
         PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


def small_label(slide, value, x, y, w, color=TEAL):
    rect(slide, x, y + 0.03, 0.04, 0.22, color, None)
    text(slide, value, x + 0.12, y, w - 0.12, 0.27, 9.4, True, color)


def table(slide, headers, rows, x, y, w, h, widths, size=8.5,
          highlight_rows=(), align_last_left=True):
    tbl = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    for i, value in enumerate(widths):
        tbl.columns[i].width = Inches(value)
    for j, value in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.04)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = value
        style_run(r, size, True, WHITE)
    for i, row in enumerate(rows, 1):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            if i - 1 in highlight_rows:
                cell.fill.fore_color.rgb = PALE_GOLD
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else PALE_BLUE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.04)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            left = j == 0 or (align_last_left and j == len(row) - 1)
            p.alignment = PP_ALIGN.LEFT if left else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(value)
            style_run(r, size, i - 1 in highlight_rows and j == 0, INK)
    return tbl


def metric(slide, value, label, x, y, w, accent=BLUE):
    text(slide, value, x, y, w, 0.38, 17, True, accent, PP_ALIGN.CENTER)
    text(slide, label, x, y + 0.42, w, 0.22, 8.4, False, GRAY, PP_ALIGN.CENTER)


def hbar(slide, label, value, x, y, w, max_value=100, color=BLUE, suffix=""):
    text(slide, label, x, y - 0.02, 1.34, 0.24, 8.5, False, INK, PP_ALIGN.RIGHT)
    rect(slide, x + 1.46, y + 0.03, w - 2.18, 0.15, PALE_BLUE, None)
    if value > 0:
        rect(slide, x + 1.46, y + 0.03, (w - 2.18) * value / max_value,
             0.15, color, None)
    text(slide, f"{value:.1f}{suffix}", x + w - 0.64, y - 0.04, 0.60, 0.25,
         8.5, True, color, PP_ALIGN.RIGHT)


# Slide 1 remains unchanged from the unified concise deck.


# 2 Background and concrete task --------------------------------------------
s = setup(prs.slides[1], "复杂雷达问答本质上是组合检索与精确执行问题", "研究背景与问题场景", 2)
box(s, "“美国研制且工作在 J 波段的雷达有多少款？”",
    1.25, 1.40, 7.50, 0.60, PALE_BLUE, BLUE, 14.0, True, NAVY)
small_label(s, "知识环境", 0.68, 2.36, 1.30)
for i, (a, b) in enumerate([
    ("RadarKG", "实体 · 关系 · 属性"),
    ("叙述文本", "原理 · 结构 · 背景"),
    ("工具接口", "图检索 · 集合 · 聚合"),
]):
    y = 2.74 + i * 0.72
    box(s, a, 0.78, y, 1.16, 0.45, WHITE, BLUE, 9.8, True, NAVY)
    text(s, b, 2.08, y + 0.11, 1.45, 0.22, 8.6, False, GRAY)
line(s, 3.72, 2.48, 3.72, 5.08, MID, 0.8)
small_label(s, "可执行查询计划", 4.06, 2.36, 1.70)
box(s, "实体 / 条件识别", 4.15, 2.72, 1.38, 0.46, WHITE, BLUE, 9.4, True, NAVY)
arrow(s, 5.58, 2.87, 0.24, 0.15, BLUE)
box(s, "constraint\ncountry = 美国", 5.88, 2.64, 1.42, 0.62, WHITE, BLUE,
    8.8, True, NAVY, name=MONO)
box(s, "constraint\nband = J", 5.88, 3.56, 1.42, 0.62, WHITE, BLUE,
    8.8, True, NAVY, name=MONO)
line(s, 7.31, 2.95, 7.62, 2.95, BLUE, 1.0)
line(s, 7.31, 3.87, 7.62, 3.87, BLUE, 1.0)
line(s, 7.62, 2.95, 7.62, 3.87, BLUE, 1.0)
arrow(s, 7.62, 3.33, 0.26, 0.16, BLUE)
box(s, "intersect", 7.92, 3.18, 1.04, 0.48, PALE_TEAL, TEAL, 9.4, True, NAVY,
    name=MONO)
arrow(s, 8.30, 3.72, 0.16, 0.25, TEAL, "down")
box(s, "count", 7.92, 4.06, 1.04, 0.46, PALE_GOLD, GOLD, 9.6, True, NAVY,
    name=MONO)
text(s, "中间变量与执行结果共同决定下一步动作", 4.24, 4.90, 4.56, 0.28,
     9.8, True, TEAL, PP_ALIGN.CENTER)
takeaway(s, "复杂雷达问答的难点不在一次召回，而在多步组合检索、精确计算和状态驱动决策。")


# 3 Related work and questions ----------------------------------------------
s = setup(prs.slides[2], "现有方法的核心缺口是计划表达与长程学习", "研究现状与科学问题", 3)
table(s, ["方法范式", "主要能力", "复杂组合问答中的不足"], [
    ["RAG", "文本证据召回", "难以表达集合、计数和多步状态"],
    ["图路径 / KGQA", "关系路径推理", "集合聚合、参数约束和变量依赖有限"],
    ["ReAct / Tool Agent", "灵活工具调用", "自由动作空间下参数生成与长程决策不稳定"],
], 0.72, 1.38, 8.56, 2.36, [1.62, 2.38, 4.56], 9.2)
small_label(s, "两个递进科学问题", 0.72, 4.12, 2.15)
box(s, "Q1  策略空间如何设计？\n复杂查询如何表示为类型安全、变量依赖明确、可确定执行的策略？",
    0.80, 4.52, 4.02, 1.14, WHITE, BLUE, 10.2, True, NAVY)
box(s, "Q2  策略如何有效学习？\n如何利用计划、进展、恢复和成本反馈缓解终局奖励稀疏？",
    5.12, 4.52, 4.10, 1.14, WHITE, TEAL, 10.2, True, NAVY)
takeaway(s, "论文研究对象是多步可执行图检索策略：策略空间构造定义“学什么”，执行反馈优化解决“怎么学”。")


# 4 Native-vector method overview ------------------------------------------
s = setup(prs.slides[3], "统一方法：可执行图检索策略学习", "总体研究思路与技术路线", 4)
small_label(s, "推理执行闭环", 0.62, 1.30, 1.65)
nodes = [
    (0.72, 1.86, 1.05, "复杂问题 q", WHITE, BLUE),
    (2.05, 1.78, 1.32, "检索策略 πθ", PALE_GOLD, GOLD),
    (3.67, 1.78, 1.42, "类型化动作\n查询计划", PALE_BLUE, BLUE),
    (5.39, 1.68, 1.58, "确定性执行环境\nGraph · Set · Count\nCompare · Text", WHITE, BLUE),
    (7.28, 1.78, 1.30, "结构化观察 o_t", PALE_TEAL, TEAL),
    (8.86, 1.86, 0.66, "答案", WHITE, BLUE),
]
for x, y, w, label, fill, border in nodes:
    box(s, label, x, y, w, 0.72 if x != 5.39 else 0.92, fill, border,
        9.2, True, NAVY)
for x in (1.80, 3.40, 5.12, 7.00, 8.61):
    arrow(s, x, 2.05, 0.20, 0.14, BLUE)
text(s, "a_t ~ πθ(a_t | s_t)", 2.04, 2.58, 1.70, 0.24, 9.0, True, NAVY,
     PP_ALIGN.CENTER, name=MONO)
text(s, "s_t = {q, history, observation}", 3.82, 2.58, 2.52, 0.24,
     8.8, False, GRAY, PP_ALIGN.CENTER, name=MONO)
line(s, 7.92, 1.74, 7.92, 1.35, TEAL, 1.2)
line(s, 7.92, 1.35, 2.70, 1.35, TEAL, 1.2)
arrow(s, 2.54, 1.34, 0.16, 0.22, TEAL, "down")
text(s, "继续规划 / 参数修正 / 错误恢复", 4.58, 1.10, 2.15, 0.22,
     8.2, True, TEAL, PP_ALIGN.CENTER)

line(s, 0.70, 3.05, 9.30, 3.05, LINE, 0.8)
small_label(s, "策略学习闭环", 0.62, 3.24, 1.65)
learn = [
    (0.78, "成功轨迹", "D_sft", WHITE, BLUE),
    (2.27, "SFT 初始化", "πθ0", PALE_BLUE, BLUE),
    (3.82, "检索策略", "πθ", PALE_GOLD, GOLD),
    (5.36, "Rollout", "轨迹 τ", WHITE, BLUE),
    (6.89, "可执行反馈", "R(τ)", PALE_TEAL, TEAL),
    (8.38, "策略优化", "update πθ", PALE_GOLD, GOLD),
]
for x, a, b, fill, border in learn:
    box(s, f"{a}\n{b}", x, 3.77, 1.18, 0.70, fill, border, 9.0, True, NAVY)
for x in (1.99, 3.51, 5.05, 6.58, 8.08):
    arrow(s, x, 4.04, 0.20, 0.14, TEAL if x > 6.0 else BLUE)
line(s, 8.97, 4.50, 8.97, 4.88, GOLD, 1.1)
line(s, 8.97, 4.88, 4.40, 4.88, GOLD, 1.1)
arrow(s, 4.26, 4.66, 0.16, 0.22, GOLD, "up")
text(s, "Answer · Plan · Progress · Recovery · Cost", 6.36, 4.55, 2.00, 0.20,
     7.8, False, TEAL, PP_ALIGN.CENTER)

rect(s, 0.78, 5.30, 8.52, 0.68, WHITE, MID)
text(s, "Generalization Evaluation", 0.98, 5.46, 1.92, 0.22,
     8.8, True, GRAY, name=MONO)
text(s, "Composition Holdout", 3.18, 5.45, 1.62, 0.22, 8.6, True, NAVY,
     PP_ALIGN.CENTER, name=MONO)
text(s, "KQA Pro", 5.30, 5.45, 1.05, 0.22, 8.6, True, NAVY,
     PP_ALIGN.CENTER, name=MONO)
text(s, "MetaQA", 6.90, 5.45, 1.05, 0.22, 8.6, True, NAVY,
     PP_ALIGN.CENTER, name=MONO)
takeaway(s, "策略空间定义可表达决策，执行反馈优化同一个策略 πθ，并由组合隔离检验泛化。")


# 5 Policy-space construction ----------------------------------------------
s = setup(prs.slides[4], "策略空间构造：将复杂查询表示为类型安全的可执行计划", "统一方法  |  策略表示", 5)
small_label(s, "状态、动作与转移", 0.64, 1.30, 1.90)
box(s, "状态 s_t\n问题 + 历史 + 中间结果 + 错误", 0.82, 1.72, 1.78, 0.72,
    WHITE, BLUE, 9.5, True, NAVY)
arrow(s, 2.70, 1.98, 0.25, 0.16, BLUE)
box(s, "动作 a_t = (macro, operator, args, bindings, control)",
    3.05, 1.65, 3.94, 0.86, PALE_GOLD, GOLD, 10.3, True, NAVY, name=MONO)
arrow(s, 7.10, 1.98, 0.25, 0.16, BLUE)
box(s, "状态转移\n执行器返回结构化观察", 7.46, 1.72, 1.72, 0.72,
    WHITE, BLUE, 9.5, True, NAVY)
text(s, "macro: 路径 / 约束 / 集合 / 聚合    operator: constraint / path / intersect / count",
     2.82, 2.66, 4.60, 0.24, 8.3, False, GRAY, PP_ALIGN.CENTER)
text(s, "args: 实体 / 关系 / 属性    bindings: 中间变量    control: continue / revise / finish",
     2.72, 2.92, 4.82, 0.24, 8.3, False, GRAY, PP_ALIGN.CENTER)

small_label(s, "类型系统", 0.64, 3.28, 1.25)
types = [("Entity", "实体"), ("EntitySet", "实体集合"), ("Relation", "关系"),
         ("Scalar", "标量"), ("Evidence", "证据")]
for i, (a, b) in enumerate(types):
    x = 0.84 + i * 1.72
    box(s, f"{a}\n{b}", x, 3.64, 1.42, 0.55, WHITE, MID, 8.6, True, BLUE,
        name=MONO if i == 0 else FONT)

small_label(s, "可执行计划树", 0.64, 4.52, 1.55)
box(s, "count($s3)\nEntitySet → Scalar", 4.08, 4.46, 1.72, 0.50,
    NAVY, None, 8.7, True, WHITE, name=MONO)
box(s, "intersect($s1, $s2)\nEntitySet × EntitySet → EntitySet", 3.54, 5.16, 2.80, 0.56,
    BLUE, None, 8.4, True, WHITE, name=MONO)
box(s, "constraint\ncountry = 美国 → $s1", 0.90, 5.18, 2.16, 0.54,
    PALE_BLUE, BLUE, 8.5, True, NAVY, name=MONO)
box(s, "constraint\nband = J → $s2", 6.82, 5.18, 2.16, 0.54,
    PALE_BLUE, BLUE, 8.5, True, NAVY, name=MONO)
line(s, 2.95, 5.45, 3.53, 5.45, BLUE, 1.2)
arrow(s, 3.34, 5.37, 0.19, 0.16, BLUE)
line(s, 6.35, 5.45, 6.92, 5.45, BLUE, 1.2)
arrow(s, 6.35, 5.37, 0.19, 0.16, BLUE, "left")
line(s, 4.94, 5.15, 4.94, 4.97, BLUE, 1.2)
arrow(s, 4.86, 4.96, 0.16, 0.19, BLUE, "up")
takeaway(s, "LLM 负责语义决策，确定性执行器负责集合、路径、计数与比较等精确计算。")


# 6 Policy learning ---------------------------------------------------------
s = setup(prs.slides[5], "策略优化：利用执行反馈学习多步规划、修正与终止", "统一方法  |  策略学习", 6)
small_label(s, "SFT 初始化", 0.64, 1.30, 1.35)
box(s, "成功轨迹 D_sft", 0.78, 1.72, 1.42, 0.56, WHITE, BLUE, 9.2, True, NAVY,
    name=MONO)
arrow(s, 2.27, 1.91, 0.22, 0.15, BLUE)
box(s, "SFT 冷启动", 2.56, 1.72, 1.30, 0.56, PALE_BLUE, BLUE, 9.4, True, NAVY)
arrow(s, 3.94, 1.91, 0.22, 0.15, BLUE)
box(s, "πθ0", 4.24, 1.72, 0.72, 0.56, PALE_GOLD, GOLD, 11.0, True, NAVY,
    name=MONO)

small_label(s, "执行反馈优化闭环", 0.64, 2.60, 1.85)
loop_nodes = [
    (0.90, "Retrieval Policy\nπθ", PALE_GOLD, GOLD),
    (2.72, "Executable\nEnvironment", WHITE, BLUE),
    (4.54, "Trajectory\nτ", PALE_BLUE, BLUE),
    (6.36, "Executable\nEvaluation", PALE_TEAL, TEAL),
    (8.18, "Policy\nOptimization", PALE_GOLD, GOLD),
]
for x, label, fill, border in loop_nodes:
    box(s, label, x, 3.08, 1.34, 0.76, fill, border, 9.3, True, NAVY, name=MONO)
for x in (2.35, 4.17, 5.99, 7.81):
    arrow(s, x, 3.38, 0.22, 0.15, TEAL if x > 5.5 else BLUE)
text(s, "action", 2.28, 3.04, 0.40, 0.18, 7.3, False, GRAY, PP_ALIGN.CENTER,
     name=MONO)
text(s, "R(τ)", 7.68, 3.04, 0.40, 0.18, 7.8, True, TEAL, PP_ALIGN.CENTER,
     name=MONO)
line(s, 8.85, 3.89, 8.85, 4.18, GOLD, 1.1)
line(s, 8.85, 4.18, 1.58, 4.18, GOLD, 1.1)
arrow(s, 1.50, 3.86, 0.16, 0.30, GOLD, "up")
text(s, "update πθ", 4.80, 4.18, 1.10, 0.20, 8.0, True, GOLD,
     PP_ALIGN.CENTER, name=MONO)
text(s, "a_t ~ πθ(a_t | s_t)       τ = (s0, a0, o1, ..., aT)",
     2.40, 4.43, 5.20, 0.26, 9.0, True, NAVY, PP_ALIGN.CENTER, name=MONO)

rect(s, 0.86, 4.90, 8.30, 0.64, WHITE, MID)
text(s, "R(τ) = R_answer + λp R_plan + λg R_progress + λf R_format − λs C_step − λr C_repeat",
     1.04, 5.08, 7.94, 0.26, 9.4, True, NAVY, PP_ALIGN.CENTER, name=MONO)
text(s, "Answer 终局正确性  |  Plan 类型合法性  |  Progress 状态进展  |  Recovery 错误修复  |  Cost 步数 / token / 延迟",
     0.86, 5.68, 8.30, 0.24, 8.0, False, GRAY, PP_ALIGN.CENTER)
text(s, "课程：基础任务  →  组合任务  →  错误恢复", 2.96, 6.06, 4.10, 0.22,
     8.7, True, TEAL, PP_ALIGN.CENTER)
takeaway(s, "围绕同一策略 πθ，利用完整执行轨迹提供可计算的中间学习信号。")


# 7 Existing foundation -----------------------------------------------------
s = setup(prs.slides[6], "知识环境、问答数据与执行器已形成完整实验基础", "已有研究基础", 7)
small_label(s, "核心研究资产", 0.64, 1.28, 1.55)
metrics = [
    ("22,241", "RadarKG 边"), ("10,744", "实体"), ("1,432", "叙述文本段"),
    ("4,277", "成功轨迹"), ("5,401", "Oracle 回放"),
]
for i, (v, lab) in enumerate(metrics):
    x = 0.72 + i * 1.78
    metric(s, v, lab, x, 1.68, 1.46, BLUE if i < 3 else TEAL)
    if i < 4:
        line(s, x + 1.58, 1.68, x + 1.58, 2.28, LINE, 0.8)

small_label(s, "数据与原型验证", 0.64, 2.62, 1.70)
table(s, ["研究资产", "规模 / 结果", "用途"], [
    ["复杂问答划分", "Train / Dev / Test = 4,277 / 582 / 542", "策略训练、选择与独立测试"],
    ["类型化执行器", "10 类算子", "路径、约束、集合、聚合与图文检索"],
    ["72 道组合题原型", "计划 100%；执行 100%；答案 98.6%（71/72）", "验证计划协议与执行链路"],
    ["知识质量评估", "人工分层抽样 n=173；85.1% [75.1, 95.0]", "界定可核验知识环境质量"],
], 0.72, 3.02, 8.56, 2.78, [1.55, 3.52, 3.49], 8.7)
text(s, "状态：本页为已完成资产与原型结果；正式 SFT/RL 对照为拟开展研究。",
     0.84, 5.98, 8.32, 0.20, 8.2, False, GRAY, PP_ALIGN.CENTER)
takeaway(s, "知识环境、复杂问答、成功轨迹和确定性执行器已经形成完整实验基础。")


# 8 Diagnostic evidence -----------------------------------------------------
s = setup(prs.slides[7], "诊断实验将研究重点定位到策略表示与策略学习", "前期实验与研究动机", 8)
s.shapes.add_picture(str(RESULT_FIG), Inches(0.62), Inches(1.42), width=Inches(4.72))
table(s, ["诊断问题", "前期结果", "研究含义"], [
    ["结构化策略是否有效？", "52.7 → 60.3 → 88.6", "类型化计划对复杂组合题具有明显价值"],
    ["单步路由是否为瓶颈？", "LLM 88.6；Oracle 88.2", "一次路由已接近上限，应关注多步计划与执行"],
    ["固定策略能否跨任务迁移？", "KQA Pro Δ≈−0.2 pp；CI 跨 0", "需要可学习、可适配的检索策略"],
], 5.52, 1.43, 3.86, 3.82, [1.32, 0.96, 1.58], 7.6)
rect(s, 5.52, 5.48, 3.86, 0.54, PALE_BLUE, BLUE)
text(s, "补充观察：Qwen Planner 经成功轨迹微调 31.5 → 61.7",
     5.70, 5.62, 3.50, 0.22, 8.3, True, NAVY, PP_ALIGN.CENTER)
text(s, "RadarKG-QA-499；点为答案准确率，线为 95% bootstrap CI",
     0.76, 5.72, 4.40, 0.20, 7.8, False, GRAY, PP_ALIGN.CENTER)
takeaway(s, "这些前期结果用于定位研究问题；正式策略学习增益由 2×2 SFT/RL 对照验证。")


# 9 Experimental design -----------------------------------------------------
s = setup(prs.slides[8], "2×2 对照独立验证策略空间构造与执行反馈优化", "实验设计与公开验证", 9)
small_label(s, "核心因果对照", 0.62, 1.30, 1.60)
text(s, "Flat Policy Space", 1.72, 1.72, 1.88, 0.23, 8.7, True, GRAY,
     PP_ALIGN.CENTER, name=MONO)
text(s, "Typed Policy Space", 4.02, 1.72, 2.02, 0.23, 8.7, True, BLUE,
     PP_ALIGN.CENTER, name=MONO)
text(s, "SFT", 0.74, 2.30, 0.50, 0.22, 9.2, True, GRAY, PP_ALIGN.CENTER, name=MONO)
text(s, "RL", 0.74, 3.34, 0.50, 0.22, 9.2, True, GRAY, PP_ALIGN.CENTER, name=MONO)
box(s, "Flat-SFT\n成功轨迹模仿", 1.58, 2.08, 2.05, 0.72, PALE_BLUE, MID,
    9.7, True, NAVY, name=MONO)
box(s, "Typed-SFT\n结构化动作 + 轨迹", 4.02, 2.08, 2.05, 0.72, WHITE, BLUE,
    9.7, True, NAVY, name=MONO)
box(s, "Flat-RL\n扁平动作 + 执行反馈", 1.58, 3.10, 2.05, 0.72, PALE_BLUE, MID,
    9.7, True, NAVY, name=MONO)
box(s, "Full Method\n类型化策略 + 多粒度反馈", 4.02, 3.10, 2.05, 0.72,
    PALE_GOLD, GOLD, 9.5, True, NAVY, name=MONO)
arrow(s, 3.68, 2.34, 0.25, 0.16, BLUE)
text(s, "验证策略空间构造", 2.70, 2.86, 2.18, 0.20, 8.0, True, BLUE, PP_ALIGN.CENTER)
line(s, 5.04, 2.83, 5.04, 3.08, GOLD, 1.3)
text(s, "验证执行反馈优化", 5.18, 2.86, 1.30, 0.36, 8.0, True, GOLD)

small_label(s, "三层验证", 6.48, 1.30, 1.20)
table(s, ["数据集", "验证任务"], [
    ["RadarKG-QA", "IID / Entity Holdout / Composition Holdout / Error Recovery"],
    ["KQA Pro", "程序组合与类型化计划"],
    ["MetaQA", "1 / 2 / 3-hop 跨领域多跳"],
], 6.48, 1.70, 2.88, 2.32, [1.02, 1.86], 7.4)
rect(s, 6.50, 4.22, 2.84, 0.90, WHITE, MID)
text(s, "Composition Holdout", 6.67, 4.34, 2.50, 0.22, 8.3, True, BLUE,
     PP_ALIGN.CENTER, name=MONO)
text(s, "训练与测试采用不同算子组合 / 程序结构，检验可组合策略而非模板记忆。",
     6.70, 4.62, 2.42, 0.36, 7.7, True, NAVY, PP_ALIGN.CENTER)
rect(s, 0.76, 5.42, 8.48, 0.58, PALE_BLUE, BLUE)
text(s, "Environment Oracle  →  Policy Execution  →  Final Answer  →  Cost / Recovery",
     1.10, 5.59, 7.80, 0.24, 9.3, True, NAVY, PP_ALIGN.CENTER, name=MONO)
takeaway(s, "横向验证策略表示，纵向验证策略优化，组合隔离检验 compositional generalization。")


# 10 Contributions, chapters, schedule -------------------------------------
s = setup(prs.slides[9], "论文围绕一个统一方法形成两项递进贡献", "创新点、论文结构与研究计划", 10)
table(s, ["递进环节", "核心贡献", "解决问题", "因果验证"], [
    ["策略空间构造", "类型化可执行图检索决策建模", "动作空间过大；参数与变量依赖不稳定", "Flat-SFT → Typed-SFT"],
    ["执行反馈优化", "基于多粒度执行反馈的图检索策略学习", "终局奖励稀疏；长程信用分配困难", "Typed-SFT → Full-RL"],
], 0.64, 1.34, 8.72, 1.70, [1.35, 2.62, 2.63, 2.12], 8.4, highlight_rows=(0, 1))
small_label(s, "论文章节", 0.66, 3.34, 1.30)
chapters = [
    ("第 3 章", "知识环境与\n复杂问答任务"),
    ("第 4 章", "图检索策略\n空间构造"),
    ("第 5 章", "执行反馈驱动的\n策略学习"),
    ("第 6 章", "实验与\n组合泛化"),
]
for i, (a, b) in enumerate(chapters):
    x = 0.76 + i * 2.17
    box(s, f"{a}\n{b}", x, 3.70, 1.70, 0.78, WHITE, BLUE if i < 3 else TEAL,
        9.1, True, NAVY)
    if i < 3:
        arrow(s, x + 1.78, 4.00, 0.22, 0.16, BLUE)
small_label(s, "研究计划", 0.66, 4.82, 1.30)
line(s, 1.02, 5.54, 8.92, 5.54, MID, 1.2)
phases = [
    (1.02, "1–2 月", "任务与数据冻结"),
    (3.58, "3–4 月", "环境与基线"),
    (6.10, "5–7 月", "SFT / RL / 消融"),
    (8.90, "8–10 月", "公开验证 / 写作"),
]
for i, (x, period, label) in enumerate(phases):
    rect(s, x - 0.06, 5.45, 0.18, 0.18, GOLD if i == 2 else BLUE, None,
         MSO_SHAPE.OVAL)
    text(s, period, x - 0.38, 5.10, 0.86, 0.22, 8.2, True,
         GOLD if i == 2 else TEAL, PP_ALIGN.CENTER)
    text(s, label, x - 0.72, 5.80, 1.48, 0.25, 8.0, True, NAVY, PP_ALIGN.CENTER)
takeaway(s, "构造可执行的检索策略空间，利用执行反馈学习多步决策，并通过组合隔离验证策略泛化。")


# A1 Retrieval ablation -----------------------------------------------------
s = new_backup("已有检索消融：图扩展提升小规模原型的知识覆盖", "备份材料  |  检索消融", 11)
small_label(s, "21 题原型准确率（%）", 0.70, 1.34, 2.05)
for i, (lab, val, color) in enumerate([
    ("Vector-only", 52.4, BLUE), ("BM25-only", 61.9, TEAL),
    ("BM25 + Vector RRF", 66.7, BLUE), ("Full + Graph", 95.2, GOLD),
]):
    hbar(s, lab, val, 0.90, 1.90 + i * 0.72, 5.30, 100, color, "%")
table(s, ["阶段", "作用"], [
    ["BM25 / Vector", "精确术语与语义证据双路召回"],
    ["RRF", "无标度依赖的排序融合"],
    ["Graph Expansion", "补充关系邻域与结构化事实"],
], 6.30, 1.72, 2.92, 2.60, [1.24, 1.68], 8.2)
text(s, "数据来源：evaluation/final_results_summary.json",
     0.92, 5.24, 4.90, 0.20, 8.0, False, GRAY)
takeaway(s, "该结果用于说明知识环境中的检索互补性；样本规模较小，正式结论以后续主实验为准。")


# A2 Planner and operator ablation -----------------------------------------
s = new_backup("Planner 微调与关键算子消融定位策略能力来源", "备份材料  |  决策机制诊断", 12)
small_label(s, "Planner / 类型化策略（RadarKG-QA-499）", 0.66, 1.30, 3.45)
for i, (lab, val, color) in enumerate([
    ("Qwen-ZS", 31.5, MID), ("Qwen-FT", 61.7, TEAL),
    ("LLaMA-ZS", 26.5, MID), ("LLaMA-FT", 55.9, TEAL),
    ("类型化策略", 88.6, GOLD),
]):
    hbar(s, lab, val, 0.72, 1.82 + i * 0.58, 4.62, 100, color, "%")
small_label(s, "关键算子移除（n=100）", 5.56, 1.30, 2.40)
table(s, ["移除算子", "准确率下降", "诊断"], [
    ["exhaustive", "−18 pp", "枚举覆盖是主要贡献"],
    ["constrained_join", "−8 pp", "多约束连接具有独立价值"],
    ["path_plan", "−6 pp", "路径规划贡献稳定"],
    ["complement", "0 pp", "当前子集覆盖不足"],
    ["dual_subgraph", "0 pp", "当前子集覆盖不足"],
], 5.54, 1.72, 3.76, 3.60, [1.38, 0.86, 1.52], 7.6)
text(s, "来源：experiments/rog_finetune/RESULTS.md；results/ablation_strategies_100_summary.md",
     0.76, 5.52, 8.44, 0.20, 7.7, False, GRAY, PP_ALIGN.CENTER)
takeaway(s, "成功轨迹改善规划器，类型化计划与关键组合算子进一步提升复杂问题执行能力。")


# A3 Reproducibility --------------------------------------------------------
s = new_backup("正式实验采用分层报告、重复运行与配对统计检验", "备份材料  |  实验统计与可复现性", 13)
table(s, ["实验环节", "正式协议", "目的"], [
    ["重复运行", "核心方法至少 3 个随机种子", "报告均值、方差与训练稳定性"],
    ["显著性检验", "题级 paired bootstrap / 随机化检验", "控制同题比较并报告置信区间"],
    ["数据划分", "IID / Entity / Composition Holdout", "区分同分布、实体和程序组合泛化"],
    ["训练子集", "仅使用可闭世界验证样本", "避免 KG 缺失被误作策略负奖励"],
    ["分层报告", "Environment Oracle → Policy Execution → Final Answer", "区分环境覆盖、策略执行与答案生成"],
    ["效率与恢复", "调用步数、token、延迟、重复率、恢复成功率", "评估策略成本和错误修正能力"],
], 0.68, 1.42, 8.64, 4.50, [1.55, 3.88, 3.21], 8.2)
rect(s, 1.02, 6.10, 7.96, 0.38, PALE_BLUE, BLUE)
text(s, "按题型、计划深度、算子数量、未见实体与环境覆盖情况分组报告。",
     1.20, 6.19, 7.60, 0.20, 8.8, True, NAVY, PP_ALIGN.CENTER)


# A4 Risks and boundaries ---------------------------------------------------
s = new_backup("主要风险通过环境分层、课程训练和受控预算处理", "备份材料  |  风险与边界", 14)
table(s, ["风险", "可能影响", "处理方案"], [
    ["知识图谱缺失或冲突", "奖励噪声与错误归因", "闭世界训练子集；报告环境 Oracle；保留来源元数据"],
    ["公开数据工具语义不一致", "跨域比较失真", "为 KQA Pro / MetaQA 重建等价执行接口并先测 Oracle"],
    ["长轨迹奖励仍然稀疏", "策略训练不稳定", "SFT 冷启动；基础→组合→恢复课程；逐项反馈消融"],
    ["策略探索带来高成本", "训练时间与调用开销增加", "限制最大步数；缓存确定性执行；统计 token 与延迟"],
    ["模型记忆程序模板", "组合泛化被高估", "实体隔离与 Composition Holdout；错误恢复扰动测试"],
], 0.68, 1.44, 8.64, 4.42, [1.80, 2.28, 4.56], 8.4)
takeaway(s, "论文结论将分别报告知识环境上限、策略执行能力和最终答案效果，避免跨层归因。")


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUTPUT))
print(f"saved: {OUTPUT}")
print(f"slides: {len(prs.slides)}; main=10; backup={len(prs.slides) - 10}")
