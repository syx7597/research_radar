# -*- coding: utf-8 -*-
"""Polish slides 4, 6, and 8 of the final thesis proposal deck."""

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "reports" / "硕士学位论文开题答辩_图检索策略学习_最终版.pptx"
OUTPUT = ROOT / "reports" / "硕士学位论文开题答辩_图检索策略学习_定稿优化版.pptx"
RESULT_FIG = ROOT / "reports" / "academic_figures" / "representative_results.png"

NAVY = RGBColor(0x18, 0x36, 0x57)
BLUE = RGBColor(0x2F, 0x64, 0x91)
TEAL = RGBColor(0x4F, 0x7D, 0x7B)
GOLD = RGBColor(0xB4, 0x88, 0x38)
INK = RGBColor(0x24, 0x2D, 0x36)
GRAY = RGBColor(0x68, 0x73, 0x7D)
MID = RGBColor(0xB5, 0xC1, 0xCC)
LINE = RGBColor(0xD8, 0xE0, 0xE7)
PALE_BLUE = RGBColor(0xEC, 0xF1, 0xF5)
PALE_TEAL = RGBColor(0xEB, 0xF2, 0xF1)
PALE_GOLD = RGBColor(0xF6, 0xF1, 0xE6)
BG = RGBColor(0xFA, 0xFB, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"
MONO = "Consolas"


prs = Presentation(str(SOURCE))
logo_shape = next(sh for sh in prs.slides[0].shapes if sh.shape_type == 13)
LOGO = BytesIO(logo_shape.image.blob)


def set_typeface(run, name=FONT):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", name)


def style_run(run, size=12, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_typeface(run, name)


def text(slide, value, x, y, w, h, size=12, bold=False, color=INK,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03, name=FONT):
    obj = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = obj.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    style_run(r, size, bold, color, name)
    return obj


def rect(slide, x, y, w, h, fill=WHITE, border=LINE,
         kind=MSO_SHAPE.RECTANGLE, width=0.8):
    obj = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = fill
    if border is None:
        obj.line.fill.background()
    else:
        obj.line.color.rgb = border
        obj.line.width = Pt(width)
    obj.shadow.inherit = False
    return obj


def box(slide, value, x, y, w, h, fill=WHITE, border=LINE, size=9.5,
        bold=False, color=INK, align=PP_ALIGN.CENTER, width=0.8, name=FONT):
    obj = rect(slide, x, y, w, h, fill, border, width=width)
    tf = obj.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.025)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    style_run(r, size, bold, color, name)
    return obj


def line(slide, x1, y1, x2, y2, color=MID, width=1.0):
    obj = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    obj.line.color.rgb = color
    obj.line.width = Pt(width)
    return obj


def arrow(slide, x, y, w, h, color=BLUE, direction="right"):
    kinds = {
        "right": MSO_SHAPE.RIGHT_ARROW,
        "left": MSO_SHAPE.LEFT_ARROW,
        "down": MSO_SHAPE.DOWN_ARROW,
        "up": MSO_SHAPE.UP_ARROW,
    }
    return rect(slide, x, y, w, h, color, None, kinds[direction])


def clear_slide(slide):
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)


def footer(slide, page):
    line(slide, 0.45, 7.13, 9.56, 7.13, LINE, 0.6)
    text(slide, "硕士学位论文开题答辩  |  图检索策略学习", 0.47, 7.18, 5.8, 0.17,
         7.8, color=GRAY)
    text(slide, f"{page:02d}", 9.08, 7.17, 0.38, 0.18, 8.0, True, GRAY,
         PP_ALIGN.RIGHT)


def setup(slide, title_value, section, page):
    clear_slide(slide)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    rect(slide, 0, 0, 0.07, 7.5, NAVY, None)
    text(slide, section, 0.47, 0.18, 4.4, 0.22, 8.8, True, TEAL)
    text(slide, title_value, 0.47, 0.48, 8.30, 0.50, 20, True, NAVY)
    line(slide, 0.47, 1.08, 9.55, 1.08, MID, 0.8)
    LOGO.seek(0)
    slide.shapes.add_picture(LOGO, Inches(7.75), Inches(0.12), width=Inches(2.03))
    footer(slide, page)
    return slide


def small_label(slide, value, x, y, w, color=TEAL):
    rect(slide, x, y + 0.03, 0.04, 0.22, color, None)
    text(slide, value, x + 0.12, y, w - 0.12, 0.27, 9.2, True, color)


def takeaway(slide, value, y=6.48):
    rect(slide, 0.64, y, 0.045, 0.34, TEAL, None)
    text(slide, value, 0.83, y + 0.02, 8.36, 0.28, 10.1, True, NAVY,
         PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


# Slide 4: Method Overview --------------------------------------------------
s = setup(prs.slides[3], "统一方法：可执行图检索策略学习", "METHOD OVERVIEW  |  总体研究思路", 4)
small_label(s, "A  推理执行闭环", 0.62, 1.23, 1.78)
infer = [
    (0.68, 1.70, 0.82, "复杂问题\nq", WHITE, BLUE),
    (1.76, 1.65, 1.08, "检索策略\nπθ", PALE_GOLD, GOLD),
    (3.10, 1.65, 1.18, "类型化\n查询计划", PALE_BLUE, BLUE),
    (4.54, 1.58, 1.48, "确定性执行环境\nGraph Retrieval · Set\nCount / Compare · Text", WHITE, BLUE),
    (6.28, 1.65, 1.12, "结构化观察\no_t", PALE_TEAL, TEAL),
    (7.66, 1.58, 1.12, "控制\n继续 · 修正\n恢复 · 终止", WHITE, TEAL),
    (9.04, 1.70, 0.48, "答案", WHITE, BLUE),
]
for x, y, w, label, fill, border in infer:
    box(s, label, x, y, w, 0.72 if x != 4.54 else 0.86, fill, border,
        8.5, True, NAVY)
for x in (1.53, 2.87, 4.31, 6.05, 7.43, 8.81):
    arrow(s, x, 1.94, 0.19, 0.14, BLUE if x < 7.4 or x > 8.8 else TEAL)
text(s, "a_t ~ πθ(a_t | s_t)", 1.54, 2.47, 1.58, 0.22, 8.1, True, NAVY,
     PP_ALIGN.CENTER, name=MONO)
text(s, "s_t = {q, history, observation}", 3.34, 2.47, 2.50, 0.22,
     7.9, False, GRAY, PP_ALIGN.CENTER, name=MONO)
line(s, 8.21, 1.55, 8.21, 1.30, TEAL, 1.1)
line(s, 8.21, 1.30, 2.30, 1.30, TEAL, 1.1)
arrow(s, 2.22, 1.30, 0.16, 0.21, TEAL, "down")
text(s, "继续规划 / 参数修正 / 错误恢复", 4.28, 1.07, 2.12, 0.20,
     7.8, True, TEAL, PP_ALIGN.CENTER)

line(s, 0.68, 2.86, 9.32, 2.86, LINE, 0.8)
small_label(s, "B  策略学习闭环", 0.62, 3.00, 1.78)
learn = [
    (0.68, "成功轨迹\nD_sft", WHITE, BLUE),
    (1.74, "SFT\nInitialization", PALE_BLUE, BLUE),
    (2.80, "Initial Policy\nπθ0", WHITE, BLUE),
    (3.86, "Retrieval Policy\nπθ", PALE_GOLD, GOLD),
    (4.92, "Rollout\nTrajectory τ", WHITE, BLUE),
    (5.98, "Executable\nFeedback R(τ)", PALE_TEAL, TEAL),
    (7.04, "Policy\nOptimization", PALE_GOLD, GOLD),
]
for x, label, fill, border in learn:
    box(s, label, x, 3.50, 0.90, 0.66, fill, border, 7.9, True, NAVY,
        name=MONO if "Policy" in label or "Trajectory" in label else FONT)
for x in (1.60, 2.66, 3.72, 4.78, 5.84, 6.90):
    arrow(s, x, 3.75, 0.12, 0.13, TEAL if x > 5.7 else BLUE)
line(s, 7.49, 4.18, 7.49, 4.49, GOLD, 1.0)
line(s, 7.49, 4.49, 4.31, 4.49, GOLD, 1.0)
arrow(s, 4.23, 4.16, 0.16, 0.31, GOLD, "up")
text(s, "update πθ", 5.42, 4.47, 0.95, 0.18, 7.5, True, GOLD,
     PP_ALIGN.CENTER, name=MONO)
text(s, "Answer · Plan · Progress · Recovery · Cost", 5.68, 4.19, 1.78, 0.18,
     7.0, False, TEAL, PP_ALIGN.CENTER)

rect(s, 8.24, 3.28, 1.16, 1.38, WHITE, MID)
text(s, "Generalization\nEvaluation", 8.34, 3.42, 0.96, 0.36,
     7.7, True, GRAY, PP_ALIGN.CENTER, name=MONO)
line(s, 8.42, 3.88, 9.22, 3.88, LINE, 0.7)
text(s, "Composition Holdout\nKQA Pro\nMetaQA", 8.35, 4.00, 0.94, 0.52,
     7.2, True, NAVY, PP_ALIGN.CENTER, name=MONO)

text(s, "ZOOM-IN", 4.35, 4.92, 1.30, 0.20, 7.4, True, GRAY,
     PP_ALIGN.CENTER, name=MONO)
line(s, 5.00, 5.11, 5.00, 5.30, MID, 0.9)
line(s, 3.60, 5.30, 6.40, 5.30, MID, 0.9)
line(s, 3.60, 5.30, 3.60, 5.45, MID, 0.9)
line(s, 6.40, 5.30, 6.40, 5.45, MID, 0.9)
box(s, "Policy Representation\n第 5 页 · 策略如何表示", 2.52, 5.46, 2.16, 0.50,
    WHITE, BLUE, 8.2, True, NAVY, name=MONO)
box(s, "Policy Learning\n第 6 页 · 策略如何学习", 5.32, 5.46, 2.16, 0.50,
    WHITE, TEAL, 8.2, True, NAVY, name=MONO)
takeaway(s, "策略空间定义可表达决策，执行反馈优化同一策略 πθ，并通过组合隔离验证泛化。")


# Slide 6: Policy Learning --------------------------------------------------
s = setup(prs.slides[5], "策略优化：利用执行反馈学习多步规划、修正与终止", "POLICY LEARNING  |  策略优化", 6)
small_label(s, "SFT Initialization", 0.62, 1.24, 1.72)
box(s, "成功轨迹\nD_sft", 0.72, 1.65, 1.02, 0.58, WHITE, BLUE, 8.7, True, NAVY,
    name=MONO)
arrow(s, 1.78, 1.86, 0.18, 0.14, BLUE)
box(s, "SFT\nInitialization", 2.00, 1.65, 1.14, 0.58, PALE_BLUE, BLUE,
    8.3, True, NAVY, name=MONO)
arrow(s, 3.18, 1.86, 0.18, 0.14, BLUE)
box(s, "Initial Policy\nπθ0", 3.40, 1.65, 1.06, 0.58, WHITE, BLUE,
    8.4, True, NAVY, name=MONO)
line(s, 3.93, 2.24, 3.93, 2.88, GOLD, 1.0)
arrow(s, 3.85, 2.76, 0.16, 0.20, GOLD, "down")

small_label(s, "Policy Learning Loop", 0.62, 2.54, 1.82)
box(s, "Retrieval Policy\nπθ", 3.20, 2.98, 1.30, 0.68, PALE_GOLD, GOLD,
    9.0, True, NAVY, name=MONO)
box(s, "Executable\nEnvironment", 5.02, 2.98, 1.30, 0.68, WHITE, BLUE,
    8.8, True, NAVY, name=MONO)
box(s, "Trajectory\nτ", 6.84, 2.98, 1.30, 0.68, PALE_BLUE, BLUE,
    9.0, True, NAVY, name=MONO)
box(s, "Executable\nEvaluation", 6.84, 4.08, 1.30, 0.68, PALE_TEAL, TEAL,
    8.8, True, NAVY, name=MONO)
box(s, "Reward\nR(τ)", 5.02, 4.08, 1.30, 0.68, WHITE, TEAL,
    9.0, True, NAVY, name=MONO)
box(s, "Policy\nOptimization", 3.20, 4.08, 1.30, 0.68, PALE_GOLD, GOLD,
    8.8, True, NAVY, name=MONO)
arrow(s, 4.57, 3.24, 0.36, 0.16, BLUE)
arrow(s, 6.39, 3.24, 0.36, 0.16, BLUE)
arrow(s, 7.37, 3.70, 0.20, 0.30, TEAL, "down")
arrow(s, 6.39, 4.33, 0.36, 0.16, TEAL, "left")
arrow(s, 4.57, 4.33, 0.36, 0.16, GOLD, "left")
line(s, 3.85, 4.80, 3.85, 5.03, GOLD, 1.0)
line(s, 3.85, 5.03, 2.74, 5.03, GOLD, 1.0)
line(s, 2.74, 5.03, 2.74, 3.33, GOLD, 1.0)
arrow(s, 2.74, 3.25, 0.38, 0.14, GOLD, "right")
text(s, "update πθ", 2.78, 4.74, 0.92, 0.18, 7.6, True, GOLD,
     PP_ALIGN.CENTER, name=MONO)
text(s, "a_t ~ πθ(a_t | s_t)", 4.48, 2.86, 1.54, 0.20, 7.7, True, NAVY,
     PP_ALIGN.CENTER, name=MONO)
text(s, "o_{t+1}", 6.34, 2.86, 0.52, 0.20, 7.7, True, TEAL,
     PP_ALIGN.CENTER, name=MONO)
text(s, "τ = (s0, a0, o1, ..., aT)", 6.48, 3.74, 2.02, 0.20,
     7.6, False, GRAY, PP_ALIGN.CENTER, name=MONO)

rect(s, 0.72, 3.00, 2.00, 1.76, WHITE, MID)
text(s, "Reward Decomposition", 0.88, 3.16, 1.68, 0.22, 8.0, True, GRAY,
     PP_ALIGN.CENTER, name=MONO)
text(s, "R(τ) = R_answer + λp R_plan\n+ λg R_progress + λf R_format\n− λs C_step − λr C_repeat",
     0.86, 3.50, 1.72, 0.74, 8.0, True, NAVY, PP_ALIGN.CENTER,
     MSO_ANCHOR.MIDDLE, name=MONO)
text(s, "Answer · Plan · Progress\nRecovery · Cost", 0.92, 4.31, 1.60, 0.34,
     7.0, False, TEAL, PP_ALIGN.CENTER)

rect(s, 0.76, 5.28, 8.44, 0.72, WHITE, MID)
text(s, "TRAINING CURRICULUM", 0.92, 5.42, 1.44, 0.20, 7.4, True, GRAY,
     name=MONO)
text(s, "基础任务\n单跳 / 计数", 2.60, 5.34, 1.28, 0.42, 8.0, True, NAVY,
     PP_ALIGN.CENTER)
arrow(s, 3.98, 5.53, 0.24, 0.13, BLUE)
text(s, "组合任务\n集合 / 多约束 / 比较", 4.34, 5.34, 1.54, 0.42,
     8.0, True, NAVY, PP_ALIGN.CENTER)
arrow(s, 5.98, 5.53, 0.24, 0.13, TEAL)
text(s, "错误恢复\n空结果 / 别名 / 非法关系", 6.34, 5.34, 1.78, 0.42,
     8.0, True, NAVY, PP_ALIGN.CENTER)
text(s, "Answer：终局正确性  |  Plan：可执行性 / 类型合法性  |  Progress：支持事实 / 状态进展  |  Recovery：错误修复  |  Cost：步数 / token / 延迟",
     0.78, 6.12, 8.40, 0.22, 7.5, False, GRAY, PP_ALIGN.CENTER)
takeaway(s, "执行环境产生轨迹，轨迹形成可执行反馈，反馈持续更新同一个检索策略 πθ。")


# Slide 8: Findings instead of a dense table --------------------------------
s = setup(prs.slides[7], "诊断实验将研究重点定位到策略表示与策略学习", "前期实验与研究动机", 8)
s.shapes.add_picture(str(RESULT_FIG), Inches(0.62), Inches(1.42), width=Inches(4.72))
small_label(s, "三个诊断性发现", 5.52, 1.34, 1.72)
findings = [
    ("Finding 1", "52.7 → 60.3 → 88.6", "结构化计划具有明显价值", BLUE),
    ("Finding 2", "LLM 88.6  |  Oracle 88.2", "单步路由已非主要瓶颈，应转向多步计划执行", TEAL),
    ("Finding 3", "KQA Pro  Δ≈−0.2 pp，CI 跨 0", "固定策略具有任务依赖，需要可学习的环境适配策略", GOLD),
]
for i, (tag, result, meaning, color) in enumerate(findings):
    y = 1.80 + i * 1.22
    rect(s, 5.58, y, 3.72, 0.98, WHITE, LINE)
    text(s, tag, 5.74, y + 0.12, 0.82, 0.20, 7.8, True, color, name=MONO)
    text(s, result, 6.62, y + 0.10, 2.48, 0.22, 8.7, True, NAVY,
         PP_ALIGN.RIGHT, name=MONO)
    line(s, 5.74, y + 0.40, 9.12, y + 0.40, LINE, 0.7)
    text(s, meaning, 5.76, y + 0.52, 3.34, 0.30, 8.0, True, INK,
         PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
rect(s, 5.58, 5.53, 3.72, 0.46, PALE_BLUE, BLUE)
text(s, "补充观察：Qwen Planner 经成功轨迹微调 31.5 → 61.7",
     5.76, 5.65, 3.36, 0.20, 8.0, True, NAVY, PP_ALIGN.CENTER)
text(s, "RadarKG-QA-499；点为答案准确率，线为 95% bootstrap CI",
     0.76, 5.72, 4.40, 0.20, 7.8, False, GRAY, PP_ALIGN.CENTER)
takeaway(s, "前期结果说明结构化计划有效，并将下一步研究定位到可学习的多步策略与环境适配。")


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUTPUT))
print(f"saved: {OUTPUT}")
print(f"slides: {len(prs.slides)}; modified: 4, 6, 8")
