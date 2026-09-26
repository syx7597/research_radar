# -*- coding: utf-8 -*-
"""Create one editable thesis-defense slide for the RadarKG build pipeline."""

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "reports" / "论文开题_图检索策略学习.pptx"
OUTPUT = ROOT / "reports" / "知识图谱构建_单页.pptx"

NAVY = RGBColor(0x18, 0x36, 0x57)
BLUE = RGBColor(0x2F, 0x64, 0x91)
TEAL = RGBColor(0x4F, 0x7D, 0x7B)
GOLD = RGBColor(0xB4, 0x88, 0x38)
INK = RGBColor(0x24, 0x2D, 0x36)
GRAY = RGBColor(0x68, 0x73, 0x7D)
MID = RGBColor(0xB5, 0xC1, 0xCC)
LINE = RGBColor(0xD8, 0xE0, 0xE7)
PALE_BLUE = RGBColor(0xEC, 0xF1, 0xF5)
PALE_TEAL = RGBColor(0xEB, 0xF2, 0xF1)
PALE_GOLD = RGBColor(0xF6, 0xF1, 0xE6)
BG = RGBColor(0xFA, 0xFB, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"
MONO = "Consolas"


prs = Presentation(str(SOURCE))
logo_shape = next(shape for shape in prs.slides[0].shapes if shape.shape_type == 13)
LOGO = BytesIO(logo_shape.image.blob)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Keep only the newly created slide while retaining the source theme and masters.
while len(prs.slides) > 1:
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])


def set_typeface(run, name=FONT):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", name)


def style_run(run, size=12, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_typeface(run, name)


def add_text(value, x, y, w, h, size=12, bold=False, color=INK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03,
             name=FONT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    style_run(run, size, bold, color, name)
    return shape


def add_rect(x, y, w, h, fill=WHITE, border=LINE,
             kind=MSO_SHAPE.RECTANGLE, width=0.8):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border
        shape.line.width = Pt(width)
    shape.shadow.inherit = False
    return shape


def add_box(value, x, y, w, h, fill=WHITE, border=LINE, size=9.0,
            bold=False, color=INK, align=PP_ALIGN.CENTER, width=0.8,
            name=FONT):
    shape = add_rect(x, y, w, h, fill, border, width=width)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.06)
    frame.margin_top = frame.margin_bottom = Inches(0.04)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    style_run(run, size, bold, color, name)
    return shape


def add_line(x1, y1, x2, y2, color=MID, width=1.0):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def add_arrow(x, y, w=0.19, h=0.14, color=BLUE):
    return add_rect(x, y, w, h, color, None, MSO_SHAPE.RIGHT_ARROW)


def stage_header(index, english, chinese, x, w, color=BLUE):
    add_text(f"{index:02d}", x, 1.30, 0.34, 0.22, 8.0, True, color,
             PP_ALIGN.LEFT, name=MONO)
    add_text(english, x + 0.36, 1.30, w - 0.36, 0.20, 7.2, True, GRAY,
             PP_ALIGN.LEFT, name=MONO)
    add_text(chinese, x, 1.55, w, 0.30, 10.2, True, NAVY)


slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG
add_rect(0, 0, 0.07, 7.5, NAVY, None)
add_text("RESEARCH FOUNDATION  |  RADARKG CONSTRUCTION", 0.47, 0.18,
         5.4, 0.22, 8.8, True, TEAL)
add_text("知识图谱构建：从异构语料到可执行雷达知识环境", 0.47, 0.48,
         8.45, 0.50, 20, True, NAVY)
add_line(0.47, 1.08, 9.55, 1.08, MID, 0.8)
LOGO.seek(0)
slide.shapes.add_picture(LOGO, Inches(7.75), Inches(0.12), width=Inches(2.03))

# Five-stage construction pipeline.
stages = [
    ("INPUT", "异构语料接入", 0.55, 1.55, BLUE),
    ("PARSING", "解析与切分", 2.24, 1.42, BLUE),
    ("EXTRACTION", "双路知识抽取", 3.81, 1.54, GOLD),
    ("FUSION & QA", "规范化融合质检", 5.50, 1.70, TEAL),
    ("KNOWLEDGE ENV.", "图文知识环境", 7.35, 2.10, BLUE),
]
for idx, (eng, chn, x, w, color) in enumerate(stages, 1):
    stage_header(idx, eng, chn, x, w, color)

add_box("百科 / 专业网站\n公开装备指南（WEG）\n机载、海用扫描手册\nWikidata 结构化数据",
        0.55, 1.96, 1.55, 1.68, WHITE, BLUE, 8.4, False, INK)
add_box("HTML / PDF 解析\nQwen3-VL 视觉转录\n结构块 / 叙述块切分\n来源与页码保留",
        2.24, 1.96, 1.42, 1.68, PALE_BLUE, BLUE, 8.4, False, INK)
add_box("确定性抽取\n规格表 → 关系 / 属性\n\n多智能体抽取\n实体侦察 → 关系抽取 → 忠实度批判",
        3.81, 1.96, 1.54, 1.68, PALE_GOLD, GOLD, 8.0, False, INK)
add_box("Schema 类型签名\n实体 / 别名 / 单位归一\n三元组去重与来源合并\n多源印证、冲突显式记录\n全局质量规则检测",
        5.50, 1.96, 1.70, 1.68, PALE_TEAL, TEAL, 8.0, False, INK)

# Split output emphasizes the graph-text complementarity.
add_box("RadarKG\n结构化事实层\n关系 · 属性 · 来源",
        7.35, 1.96, 0.98, 1.68, PALE_BLUE, BLUE, 8.4, True, NAVY)
add_box("叙述文本库\n语义知识层\n原理 · 结构 · 背景",
        8.47, 1.96, 0.98, 1.68, WHITE, TEAL, 8.4, True, NAVY)

for x, color in ((2.10, BLUE), (3.67, BLUE), (5.36, GOLD), (7.21, TEAL)):
    add_arrow(x, 2.72, 0.13, 0.14, color)

# Audit trail under the pipeline.
add_rect(0.55, 3.88, 6.65, 0.46, WHITE, LINE)
add_text("质量链", 0.69, 3.99, 0.55, 0.18, 8.0, True, TEAL,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
add_text("原文证据保留  →  来源分层  →  词法门控 / 忠实度裁判  →  多源印证  →  冲突与异常记录",
         1.38, 3.96, 5.58, 0.22, 8.4, True, NAVY,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
add_text("结构化事实支撑多跳、集合与精确计算；叙述文本补充原理、结构和背景知识。",
         7.38, 3.91, 2.04, 0.39, 8.0, True, TEAL,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# Final scale: only verified, current figures.
add_text("CURRENT KNOWLEDGE BASE  |  当前知识库规模", 0.55, 4.62, 3.5, 0.22,
         8.2, True, TEAL, name=MONO)
metrics = [
    ("10,744", "正式实体", BLUE),
    ("21,928", "现行可用关系边", BLUE),
    ("16,609", "属性记录", TEAL),
    ("1,432", "叙述文本段", TEAL),
]
for i, (value, label, color) in enumerate(metrics):
    x = 0.55 + i * 2.24
    add_rect(x, 5.00, 2.00, 0.94, WHITE, LINE)
    add_rect(x, 5.00, 0.045, 0.94, color, None)
    add_text(value, x + 0.16, 5.12, 1.67, 0.31, 18.5, True, NAVY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, name=MONO)
    add_text(label, x + 0.16, 5.51, 1.67, 0.20, 8.6, True, GRAY,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

add_rect(0.64, 6.32, 0.045, 0.34, TEAL, None)
add_text("形成“知识图谱 + 叙述文本”的双知识环境，为后续可执行图检索、复杂问答数据构造与策略学习提供基础。",
         0.83, 6.34, 8.45, 0.28, 10.0, True, NAVY,
         PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

# Footer is consistent with the current defense deck; page number is omitted for insertion flexibility.
add_line(0.45, 7.13, 9.56, 7.13, LINE, 0.6)
add_text("硕士学位论文开题答辩  |  图检索策略学习", 0.47, 7.18,
         5.8, 0.17, 7.8, False, GRAY)
add_text("KG", 9.08, 7.17, 0.38, 0.18, 8.0, True, GRAY, PP_ALIGN.RIGHT,
         name=MONO)

prs.core_properties.title = "知识图谱构建：从异构语料到可执行雷达知识环境"
prs.core_properties.subject = "硕士学位论文开题答辩补充单页"
prs.save(str(OUTPUT))
print(f"saved: {OUTPUT}")
print(f"slides: {len(prs.slides)}")
