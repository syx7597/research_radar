# -*- coding: utf-8 -*-
"""Generate the master's thesis proposal defense deck from PPT模板1.pptx.

Run with the minimind environment, which provides python-pptx:
  C:/Users/86188/miniconda3/envs/minimind/python.exe reports/make_thesis_proposal_ppt.py
"""

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = next(ROOT.glob("*.pptx"))
OUTPUT = ROOT / "reports" / "硕士学位论文开题答辩_图检索策略学习.pptx"

# Academic palette: navy/teal as method colors, orange for planned work, red for risk.
NAVY = RGBColor(0x18, 0x35, 0x54)
BLUE = RGBColor(0x2F, 0x5C, 0x8A)
TEAL = RGBColor(0x2F, 0x7D, 0x78)
ORANGE = RGBColor(0xD2, 0x75, 0x17)
RED = RGBColor(0xB6, 0x4B, 0x4B)
DARK = RGBColor(0x20, 0x29, 0x33)
GRAY = RGBColor(0x66, 0x70, 0x7C)
MID = RGBColor(0xA9, 0xB4, 0xC0)
LIGHT = RGBColor(0xEA, 0xF0, 0xF5)
PALE_TEAL = RGBColor(0xE8, 0xF3, 0xF1)
PALE_ORANGE = RGBColor(0xFA, 0xF0, 0xE3)
PALE_RED = RGBColor(0xF8, 0xEA, 0xE8)
PALE_BLUE = RGBColor(0xE8, 0xEF, 0xF6)
BG = RGBColor(0xF8, 0xFA, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x0D, 0x0D, 0x0D)

FONT = "微软雅黑"
MONO = "Consolas"


prs = Presentation(str(TEMPLATE))
SW, SH = prs.slide_width, prs.slide_height

# Keep the template's school asset and master/theme, replace its two starter slides.
logo_shape = next(sh for sh in prs.slides[0].shapes if sh.shape_type == 13)
LOGO = BytesIO(logo_shape.image.blob)
while len(prs.slides):
    slide_id = prs.slides._sldIdLst[0]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[0]
BLANK = prs.slide_layouts[6]

prs.core_properties.title = "面向复杂雷达情报问答的图检索策略学习方法研究"
prs.core_properties.subject = "硕士学位论文开题答辩"
prs.core_properties.author = ""
prs.core_properties.keywords = "雷达情报问答, 图检索, 策略学习, 类型化执行, SFT, RL"


def _set_typeface(run, name=FONT):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", name)


def style_run(run, size=14, bold=False, color=DARK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _set_typeface(run, name)


def add_text(slide, text, x, y, w, h, size=14, bold=False, color=DARK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, name=FONT,
             margin=0.04, line_spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    style_run(run, size, bold, color, name)
    return tb


def add_rich_text(slide, spans, x, y, w, h, size=14, align=PP_ALIGN.LEFT,
                  valign=MSO_ANCHOR.TOP, margin=0.04):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    for text, bold, color, font_size in spans:
        run = p.add_run(); run.text = text
        style_run(run, font_size or size, bold, color)
    return tb


def add_bullets(slide, items, x, y, w, h, size=14, color=DARK,
                bullet_color=BLUE, gap=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.left_margin = Inches(0.15 + 0.22 * level)
        p.first_line_indent = Inches(-0.13)
        bullet = "•" if level == 0 else "–"
        rb = p.add_run(); rb.text = bullet + " "
        style_run(rb, size - level, True, bullet_color if level == 0 else GRAY)
        rt = p.add_run(); rt.text = text
        style_run(rt, size - level, False, color if level == 0 else GRAY)
    return tb


def add_shape(slide, x, y, w, h, fill=WHITE, line=LIGHT,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_width=1.0):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_width)
    sp.shadow.inherit = False
    return sp


def add_box_text(slide, text, x, y, w, h, fill=WHITE, line=LIGHT,
                 size=13, bold=False, color=DARK, align=PP_ALIGN.CENTER,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_width=1.0,
                 margin=0.08):
    sp = add_shape(slide, x, y, w, h, fill, line, shape, line_width)
    tf = sp.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; style_run(r, size, bold, color)
    return sp


def add_line(slide, x1, y1, x2, y2, color=MID, width=1.5, dash=None):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color; line.line.width = Pt(width)
    if dash is not None:
        line.line.dash_style = dash
    return line


def add_arrow(slide, x, y, w, h, color=MID, direction="right"):
    shape = {
        "right": MSO_SHAPE.RIGHT_ARROW,
        "left": MSO_SHAPE.LEFT_ARROW,
        "down": MSO_SHAPE.DOWN_ARROW,
        "up": MSO_SHAPE.UP_ARROW,
    }[direction]
    return add_shape(slide, x, y, w, h, color, None, shape)


def add_pill(slide, text, x, y, w, fill=PALE_BLUE, color=BLUE, size=10.5):
    return add_box_text(slide, text, x, y, w, 0.34, fill, None, size, True, color)


def add_status(slide, text, fill, x=8.28, y=0.72, w=1.15):
    add_pill(slide, text, x, y, w, fill, WHITE, 10)


def new_slide(title, section, page, status=None):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    add_shape(slide, 0, 0, 0.10, 7.5, NAVY, None, MSO_SHAPE.RECTANGLE)
    add_text(slide, section.upper(), 0.43, 0.18, 2.4, 0.26, 9.5, True, TEAL)
    add_text(slide, title, 0.43, 0.47, 7.35, 0.58, 22, True, NAVY)
    add_line(slide, 0.43, 1.10, 9.58, 1.10, MID, 1.0)
    LOGO.seek(0)
    slide.shapes.add_picture(LOGO, Inches(7.65), Inches(0.12), width=Inches(2.18))
    add_line(slide, 0.43, 7.15, 9.58, 7.15, LIGHT, 0.8)
    add_text(slide, "硕士学位论文开题答辩  ·  图检索策略学习", 0.46, 7.19,
             6.4, 0.20, 8.5, False, GRAY)
    add_text(slide, f"{page:02d}", 9.15, 7.17, 0.35, 0.22, 9, True, GRAY,
             PP_ALIGN.RIGHT)
    if status:
        text, fill = status
        add_status(slide, text, fill)
    return slide


def add_section_label(slide, text, x, y, w, color=BLUE):
    add_shape(slide, x, y + 0.03, 0.06, 0.30, color, None, MSO_SHAPE.RECTANGLE)
    add_text(slide, text, x + 0.16, y, w - 0.16, 0.36, 12, True, color)


def add_metric(slide, value, label, x, y, w=1.55, fill=WHITE, color=BLUE,
               note=None):
    add_shape(slide, x, y, w, 1.00 if note is None else 1.18, fill, LIGHT)
    add_text(slide, value, x + 0.08, y + 0.10, w - 0.16, 0.40, 22, True, color,
             PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.08, y + 0.56, w - 0.16, 0.28, 10.5, False, GRAY,
             PP_ALIGN.CENTER)
    if note:
        add_text(slide, note, x + 0.08, y + 0.84, w - 0.16, 0.22, 8.5, False, GRAY,
                 PP_ALIGN.CENTER)


def add_table(slide, headers, rows, x, y, w, h, widths=None, font_size=10.5,
              header_fill=NAVY, highlight_rows=()):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y),
                                   Inches(w), Inches(h)).table
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Inches(width)
    for j, head in enumerate(headers):
        cell = table.cell(0, j); cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
        cell.margin_left = cell.margin_right = Inches(0.03)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = head; style_run(r, font_size, True, WHITE)
    for i, row in enumerate(rows, 1):
        for j, value in enumerate(row):
            cell = table.cell(i, j); cell.fill.solid()
            cell.fill.fore_color.rgb = PALE_BLUE if i in highlight_rows else (WHITE if i % 2 else LIGHT)
            cell.margin_left = cell.margin_right = Inches(0.03)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run(); r.text = str(value)
            style_run(r, font_size, i in highlight_rows, BLUE if i in highlight_rows else DARK)
    return table


# ---------------------------------------------------------------------------
# 1. Cover
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
add_shape(s, 0, 0, 0.15, 7.5, NAVY, None, MSO_SHAPE.RECTANGLE)
LOGO.seek(0); s.shapes.add_picture(LOGO, Inches(6.95), Inches(0.12), width=Inches(2.85))

# Editable graph motif.
for x, y, r, c in [(7.25, 4.85, 0.18, BLUE), (8.15, 4.25, 0.13, TEAL),
                   (8.85, 5.10, 0.16, ORANGE), (7.75, 5.75, 0.12, NAVY),
                   (9.35, 5.85, 0.10, TEAL), (6.80, 6.10, 0.10, ORANGE)]:
    add_shape(s, x, y, r, r, c, None, MSO_SHAPE.OVAL)
for a, b in [((7.34, 4.94), (8.21, 4.31)), ((8.21, 4.31), (8.93, 5.18)),
             ((7.34, 4.94), (7.81, 5.81)), ((8.93, 5.18), (9.40, 5.90)),
             ((7.81, 5.81), (9.40, 5.90)), ((7.81, 5.81), (6.85, 6.15))]:
    add_line(s, *a, *b, LIGHT, 2.0)

add_pill(s, "硕士学位论文开题答辩", 0.75, 1.36, 2.15, PALE_TEAL, TEAL, 11)
add_text(s, "面向复杂雷达情报问答的\n图检索策略学习方法研究",
         0.74, 2.00, 8.35, 1.55, 30, True, NAVY, PP_ALIGN.LEFT,
         MSO_ANCHOR.MIDDLE)
add_text(s, "Graph Retrieval Policy Learning for Complex Radar Intelligence Question Answering",
         0.78, 3.72, 8.20, 0.50, 12.5, False, GRAY)
add_shape(s, 0.78, 4.45, 3.70, 0.04, TEAL, None, MSO_SHAPE.RECTANGLE)
add_text(s, "汇报人：XXX    指导教师：XXX\n专业：XXX      日期：2026 年  月  日",
         0.78, 4.78, 5.75, 0.90, 15, False, DARK)
add_text(s, "研究主线：类型化图检索决策  ×  多粒度可执行反馈",
         0.78, 6.52, 7.7, 0.38, 13, True, BLUE)


# 2. Background
s = new_slide("复杂雷达问答本质上是组合检索问题", "研究背景", 2)
add_section_label(s, "知识环境：异构且关系密集", 0.48, 1.32, 3.0)
sources = [("装备手册", "PDF / 扫描文档"), ("结构化库", "实体 · 关系 · 属性"),
           ("叙述文本", "原理 · 结构 · 背景")]
for i, (a, b) in enumerate(sources):
    y = 1.80 + i * 1.08
    add_box_text(s, a, 0.55, y, 1.18, 0.68, PALE_BLUE if i < 2 else PALE_TEAL,
                 BLUE if i < 2 else TEAL, 12, True, NAVY)
    add_text(s, b, 1.88, y + 0.08, 1.45, 0.50, 10.5, False, GRAY,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    add_arrow(s, 3.35, y + 0.22, 0.32, 0.18, MID)

add_box_text(s, "复杂问题", 3.82, 2.57, 1.30, 0.84, NAVY, None, 15, True, WHITE)
add_text(s, "“找出美国研制、工作在 J 波段的雷达，\n并统计其数量与最大探测距离”",
         3.48, 3.62, 2.05, 0.85, 11.5, False, DARK, PP_ALIGN.CENTER)

add_section_label(s, "所需能力：不是一次 Top-K", 5.70, 1.32, 3.6, ORANGE)
ops = [("多跳路径", "Radar → Platform → Country"), ("集合运算", "交 ∩ / 并 ∪ / 差 −"),
       ("精确计算", "Count / Aggregate / Compare"), ("图文协同", "事实主干 + 文本补充")]
for i, (a, b) in enumerate(ops):
    y = 1.76 + i * 0.93
    add_box_text(s, a, 5.82, y, 1.28, 0.62, PALE_ORANGE, ORANGE, 11.5, True, ORANGE)
    add_text(s, b, 7.25, y + 0.04, 2.1, 0.54, 10.5, False, DARK,
             PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

add_box_text(s, "研究对象", 0.60, 5.72, 1.05, 0.42, TEAL, None, 10.5, True, WHITE)
add_text(s, "自然语言问题如何转化为可执行、可反馈、可泛化的图检索策略",
         1.82, 5.69, 7.25, 0.48, 15, True, NAVY)
add_text(s, "雷达情报问答是应用场景；知识图谱与文本库是知识环境。",
         0.62, 6.40, 8.65, 0.35, 10.5, False, GRAY, PP_ALIGN.CENTER)


# 3. Gaps
s = new_slide("现有方法在复杂组合问答中存在两类结构性缺口", "问题提出", 3)
add_table(s, ["范式", "主要能力", "复杂组合问答中的局限"],
          [["传统 RAG", "一次文本 Top-K", "难以保证完整路径、集合与精确计数"],
           ["GraphRAG / KGQA", "图结构或关系路径", "多采用固定流程；集合、聚合与数据流表达有限"],
           ["工具 Agent / ReAct", "多轮调用与环境观察", "扁平工具空间大；参数错误、重复调用和误差传播"],
           ["检索策略 RL", "从答案反馈学习搜索", "多依赖最终 EM/F1；中间决策信用分配不足"]],
          0.54, 1.43, 8.92, 2.65, [1.55, 2.30, 5.07], 10.5)

add_box_text(s, "缺口 1  ·  Action-space design", 0.58, 4.43, 4.15, 0.52,
             PALE_BLUE, BLUE, 12, True, BLUE)
add_bullets(s, ["LLM 同时承担任务分解、工具选择、参数生成与精确计算",
                "动作组合随工具和步骤增长，长程决策稳定性下降"],
            0.65, 5.04, 3.98, 1.20, 11.5, bullet_color=BLUE, gap=5)

add_box_text(s, "缺口 2  ·  Credit assignment", 5.05, 4.43, 4.15, 0.52,
             PALE_ORANGE, ORANGE, 12, True, ORANGE)
add_bullets(s, ["最终答案 EM/F1 只能评价整条轨迹，奖励稀疏",
                "无法定位哪一步计划、执行或恢复行为产生贡献"],
            5.12, 5.04, 3.98, 1.20, 11.5, bullet_color=ORANGE, gap=5)
add_text(s, "相关工作：RAG (Lewis et al., 2020)；ReAct (Yao et al., 2023)；ToG / RoG (ICLR 2024)；Search-R1 / Graph-R1。",
         0.61, 6.67, 8.8, 0.26, 8.5, False, GRAY, PP_ALIGN.CENTER)


# 4. Scientific questions
s = new_slide("论文围绕“决策表示 → 策略学习”两个问题递进展开", "科学问题", 4)
add_box_text(s, "Q1", 0.65, 1.45, 0.62, 0.62, BLUE, None, 14, True, WHITE,
             shape=MSO_SHAPE.OVAL)
add_text(s, "复杂图检索决策如何被可执行地表示？", 1.46, 1.43, 3.35, 0.58,
         17, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
add_box_text(s, "类型化 / 分层决策建模", 0.75, 2.25, 3.85, 0.55,
             PALE_BLUE, BLUE, 13, True, BLUE)
add_bullets(s, ["语义宏动作与结构化参数分离", "类型约束、中间变量和数据依赖",
                "确定性执行集合、路径、计数与比较"],
            0.83, 2.98, 3.68, 1.75, 12, bullet_color=BLUE, gap=7)

add_arrow(s, 4.67, 3.22, 0.58, 0.34, TEAL)

add_box_text(s, "Q2", 5.38, 1.45, 0.62, 0.62, ORANGE, None, 14, True, WHITE,
             shape=MSO_SHAPE.OVAL)
add_text(s, "长程检索策略如何获得有效学习信号？", 6.17, 1.43, 3.25, 0.58,
         17, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
add_box_text(s, "多粒度可执行反馈", 5.48, 2.25, 3.80, 0.55,
             PALE_ORANGE, ORANGE, 13, True, ORANGE)
add_bullets(s, ["计划合法性与变量依赖可自动校验", "支持事实、状态进展与错误修复可观测",
                "正确率与调用成本联合优化"],
            5.56, 2.98, 3.63, 1.75, 12, bullet_color=ORANGE, gap=7)

add_box_text(s, "总体目标", 0.82, 5.22, 1.12, 0.46, TEAL, None, 11.5, True, WHITE)
add_text(s, "学习可执行、经济且具有组合泛化能力的图检索策略",
         2.15, 5.16, 6.85, 0.58, 17, True, NAVY, PP_ALIGN.CENTER,
         MSO_ANCHOR.MIDDLE)
add_text(s, "组合隔离是核心证据：检验模型学习的是策略，而不是训练模板。",
         1.05, 6.21, 7.95, 0.42, 12, False, GRAY, PP_ALIGN.CENTER)


# 5. Overall technical route
s = new_slide("从复杂问题到策略学习：总体技术路线", "总体框架", 5,
              ("核心页", TEAL))
# Main inference lane.
lane_y = 1.52
stages = [
    ("自然语言\n复杂问题", 0.42, 0.98, NAVY),
    ("语义 / 策略\n决策", 1.62, 1.05, BLUE),
    ("类型化\n查询计划", 2.90, 1.03, TEAL),
    ("确定性工具\n图查询 · 文本检索", 4.16, 1.22, NAVY),
    ("结构化观察\n结果 · 状态 · 错误", 5.62, 1.15, TEAL),
    ("继续 / 修正\n或终止", 6.99, 1.06, ORANGE),
]
for text, x, w, color in stages:
    add_box_text(s, text, x, lane_y, w, 0.90, color, None, 11.3, True, WHITE)
for x in [1.42, 2.69, 3.95, 5.40, 6.79]:
    add_arrow(s, x, lane_y + 0.32, 0.20, 0.20, MID)

# Replan loop and final answer.
add_line(s, 7.52, 2.50, 7.52, 3.05, ORANGE, 1.8)
add_line(s, 7.52, 3.05, 2.18, 3.05, ORANGE, 1.8)
add_arrow(s, 1.82, 2.91, 0.36, 0.24, ORANGE, "left")
add_text(s, "观察驱动继续规划 / 修正", 3.82, 2.74, 2.15, 0.26, 9.5, True, ORANGE,
         PP_ALIGN.CENTER)
add_arrow(s, 8.08, lane_y + 0.31, 0.20, 0.22, TEAL)
add_box_text(s, "最终答案\n结构化核心 + 受约束生成", 8.32, 1.47, 1.43, 1.02,
             PALE_TEAL, TEAL, 9.3, True, TEAL)

# Training side.
add_section_label(s, "训练侧：执行环境提供多粒度反馈", 0.47, 3.42, 4.2, ORANGE)
feedbacks = [("Answer", BLUE), ("Plan", TEAL), ("Progress", ORANGE),
             ("Recovery", RED), ("Cost", GRAY)]
for i, (txt, color) in enumerate(feedbacks):
    add_pill(s, txt, 0.63 + i * 1.17, 3.96, 1.00, color, WHITE, 9.2)
add_text(s, "执行环境", 6.70, 3.98, 1.03, 0.34, 10, True, GRAY,
         PP_ALIGN.CENTER)
add_arrow(s, 6.18, 4.03, 0.38, 0.18, ORANGE, "left")
add_line(s, 6.20, 3.88, 6.20, 2.55, ORANGE, 1.6)
add_arrow(s, 6.07, 2.48, 0.25, 0.30, ORANGE, "up")

add_section_label(s, "策略优化流程", 0.47, 4.60, 2.7, BLUE)
train = [("成功轨迹", "4,277 条已生成", PALE_BLUE, BLUE),
         ("SFT 冷启动", "学习动作协议", PALE_TEAL, TEAL),
         ("策略优化", "利用执行反馈", PALE_ORANGE, ORANGE)]
for i, (a, b, fill, color) in enumerate(train):
    x = 0.60 + i * 2.15
    add_box_text(s, a + "\n" + b, x, 5.13, 1.78, 0.80, fill, color, 10.5, True, color)
    if i < 2:
        add_arrow(s, x + 1.84, 5.42, 0.24, 0.18, MID)
add_box_text(s, "检索策略 πθ", 7.23, 5.12, 1.40, 0.82, NAVY, None, 12, True, WHITE)
add_arrow(s, 6.98, 5.41, 0.20, 0.20, BLUE)
add_line(s, 7.94, 5.08, 7.94, 4.42, ORANGE, 1.5)
add_arrow(s, 7.82, 4.32, 0.23, 0.28, ORANGE, "up")
add_text(s, "Agent 是交互形式；SFT / RL 是优化手段；研究核心是决策表示与反馈设计。",
         0.62, 6.47, 8.72, 0.35, 10.5, False, GRAY, PP_ALIGN.CENTER)


# 6. Method 1 comparison
s = new_slide("类型化决策降低复杂图检索的动作空间", "方法一 · Action-space design", 6,
              ("拟研究", ORANGE))
# Flat side.
add_box_text(s, "扁平工具 Agent", 0.55, 1.40, 4.15, 0.48, PALE_RED, RED, 13, True, RED)
add_text(s, "LLM 每一步直接面对工具 × 参数 × 历史状态", 0.70, 2.00, 3.85, 0.36,
         11, False, GRAY, PP_ALIGN.CENTER)
tools = ["graph_lookup", "subgraph", "set_op", "count", "attr_filter", "text_search"]
for i, tool in enumerate(tools):
    x = 0.72 + (i % 3) * 1.22; y = 2.52 + (i // 3) * 0.70
    add_box_text(s, tool, x, y, 1.05, 0.44, WHITE, RED, 8.6, False, RED)
add_text(s, "+ 自由参数组合", 1.54, 4.02, 2.13, 0.35, 12, True, RED, PP_ALIGN.CENTER)
add_arrow(s, 2.39, 4.42, 0.34, 0.26, RED, "down")
add_box_text(s, "模型还需自行计算集合 / 数值", 1.15, 4.82, 3.00, 0.52,
             PALE_RED, RED, 11, True, RED)
add_bullets(s, ["参数易失配", "结果易漂移", "错误传播难定位"],
            0.92, 5.55, 3.4, 0.95, 10.5, bullet_color=RED, gap=4)

# Typed side.
add_box_text(s, "本文：类型化 / 分层决策", 5.05, 1.40, 4.35, 0.48,
             PALE_TEAL, TEAL, 13, True, TEAL)
right = [("1  选择语义宏动作", "路径 / 约束 / 集合 / 聚合 / 文本补充"),
         ("2  生成类型安全参数", "Entity · Relation · EntitySet · Scalar"),
         ("3  确定性执行计算", "交并差 · 路径 · 计数 · 比较")]
for i, (a, b) in enumerate(right):
    y = 2.15 + i * 1.12
    add_box_text(s, a, 5.23, y, 1.78, 0.50, TEAL if i < 2 else NAVY,
                 None, 10.5, True, WHITE)
    add_arrow(s, 7.09, y + 0.15, 0.27, 0.18, MID)
    add_box_text(s, b, 7.44, y, 1.70, 0.50, WHITE, TEAL, 9.2, False, DARK)
add_line(s, 6.12, 2.68, 6.12, 4.33, TEAL, 1.5)
add_arrow(s, 6.01, 4.24, 0.22, 0.28, TEAL, "down")
add_box_text(s, "结构化观察返回策略", 5.72, 5.50, 2.85, 0.55,
             PALE_TEAL, TEAL, 11.5, True, TEAL)
add_text(s, "低层是确定性执行器，不是独立学习策略；因此不称为经典层次化强化学习。",
         5.12, 6.32, 4.20, 0.42, 10, False, GRAY, PP_ALIGN.CENTER)


# 7. Method 1 executable plan example
s = new_slide("类型化计划把组合问题编译为可执行数据流", "方法一 · 可执行表示", 7,
              ("拟研究", ORANGE))
add_box_text(s, "示例问题", 0.52, 1.40, 1.03, 0.42, NAVY, None, 10.5, True, WHITE)
add_text(s, "“美国研制且工作在 J 波段的雷达有多少款？”",
         1.72, 1.36, 7.42, 0.52, 15, True, NAVY, PP_ALIGN.CENTER,
         MSO_ANCHOR.MIDDLE)

# Tree.
add_box_text(s, "count\nEntitySet → Scalar", 3.65, 2.16, 2.08, 0.72,
             NAVY, None, 11.5, True, WHITE)
add_box_text(s, "intersect\nEntitySet × EntitySet → EntitySet", 3.23, 3.28, 2.92, 0.78,
             TEAL, None, 10.5, True, WHITE)
add_box_text(s, "constraint\ncountryOfOrigin = 美国", 1.15, 4.70, 2.80, 0.80,
             PALE_BLUE, BLUE, 10.5, True, BLUE)
add_box_text(s, "constraint\nhasFrequencyBand = J", 5.44, 4.70, 2.80, 0.80,
             PALE_BLUE, BLUE, 10.5, True, BLUE)
add_line(s, 4.69, 2.89, 4.69, 3.27, NAVY, 1.8)
add_line(s, 4.69, 4.07, 2.55, 4.70, TEAL, 1.8)
add_line(s, 4.69, 4.07, 6.84, 4.70, TEAL, 1.8)
add_box_text(s, "变量 $s1", 2.12, 5.82, 0.90, 0.36, PALE_TEAL, None, 9, True, TEAL)
add_box_text(s, "变量 $s2", 6.40, 5.82, 0.90, 0.36, PALE_TEAL, None, 9, True, TEAL)

# Contract and operators.
add_section_label(s, "统一执行协议", 0.52, 6.32, 2.3, BLUE)
add_text(s, "JSON Schema 校验  ·  输入输出类型  ·  中间变量  ·  数据依赖  ·  错误类别",
         2.25, 6.27, 6.85, 0.42, 10.8, False, DARK, PP_ALIGN.CENTER)
add_text(s, "已实现算子：constraint / path / intersect / union / difference / count / enumerate / aggregate / compare / contains",
         0.70, 6.72, 8.64, 0.28, 8.7, False, GRAY, PP_ALIGN.CENTER)


# 8. Method 2 reward
s = new_slide("多粒度可执行反馈缓解稀疏奖励", "方法二 · Credit assignment", 8,
              ("拟研究", ORANGE))
add_box_text(s,
             "R = R_answer + λp R_plan + λg R_progress + λf R_format − λs C_step − λr C_repeat",
             0.70, 1.43, 8.65, 0.72, NAVY, None, 15.5, True, WHITE)

rewards = [
    ("Answer", "终局 EM / F1\n计数与集合正确性", BLUE, PALE_BLUE),
    ("Plan", "计划可解析\n类型与依赖合法", TEAL, PALE_TEAL),
    ("Progress", "支持事实覆盖\n状态向目标收敛", ORANGE, PALE_ORANGE),
    ("Format", "JSON / 工具\n字段协议正确", NAVY, LIGHT),
    ("Cost", "步数 · 重复\ntoken · 延迟", RED, PALE_RED),
]
for i, (a, b, c, fill) in enumerate(rewards):
    x = 0.55 + i * 1.82
    add_box_text(s, a, x, 2.55, 1.55, 0.44, c, None, 11.5, True, WHITE)
    add_box_text(s, b, x, 3.02, 1.55, 1.02, fill, c, 9.8, False, DARK)

add_text(s, "局部信用信号", 2.55, 4.42, 1.35, 0.32, 10.5, True, TEAL,
         PP_ALIGN.CENTER)
add_arrow(s, 3.90, 4.48, 1.10, 0.20, TEAL)
add_box_text(s, "中间动作可归因", 5.15, 4.30, 1.62, 0.55,
             PALE_TEAL, TEAL, 11, True, TEAL)
add_arrow(s, 6.92, 4.48, 0.75, 0.20, ORANGE)
add_box_text(s, "长程策略可优化", 7.80, 4.30, 1.55, 0.55,
             PALE_ORANGE, ORANGE, 11, True, ORANGE)

add_bullets(s, ["错误恢复可作为 R_progress 的子项，或单独定义 R_recovery",
                "终局答案门控过程奖励，避免模型通过无意义步骤“刷分”",
                "无唯一金标计划时使用执行等价性，而非强制路径完全匹配"],
            0.80, 5.18, 8.42, 1.15, 11.5, bullet_color=ORANGE, gap=5)
add_text(s, "研究贡献在反馈构造与使用，不归结为 GRPO / PPO 等具体优化器。",
         0.75, 6.57, 8.50, 0.30, 10.5, True, GRAY, PP_ALIGN.CENTER)


# 9. Training curriculum
s = new_slide("从模仿成功轨迹到学习组合策略与错误恢复", "训练流程", 9,
              ("拟开展", ORANGE))
stages = [
    ("01", "SFT 冷启动", "4,277 条成功轨迹\n学习动作协议与基本工具调用", BLUE),
    ("02", "基础任务优化", "单跳 · 简单路径 · 计数\n答案与格式反馈", TEAL),
    ("03", "组合任务优化", "集合 · 多约束 · 比较\n计划与进展反馈", ORANGE),
    ("04", "错误恢复训练", "空结果 · 别名 · 非法关系\n修正与停止行为", RED),
]
for i, (num, title, body, color) in enumerate(stages):
    x = 0.45 + i * 2.35
    add_box_text(s, num, x + 0.67, 1.48, 0.56, 0.56, color, None, 12, True, WHITE,
                 shape=MSO_SHAPE.OVAL)
    add_box_text(s, title, x, 2.24, 1.90, 0.50, color, None, 11.2, True, WHITE)
    add_box_text(s, body, x, 2.77, 1.90, 1.18, WHITE, color, 10.2, False, DARK)
    if i < 3:
        add_arrow(s, x + 1.96, 3.20, 0.30, 0.20, MID)

add_section_label(s, "当前状态边界", 0.52, 4.35, 2.2, NAVY)
add_box_text(s, "已完成", 0.65, 4.90, 1.05, 0.43, TEAL, None, 10.5, True, WHITE)
add_text(s, "数据生成与划分、可执行工具环境、成功轨迹、SFT/策略优化脚本原型",
         1.90, 4.84, 7.25, 0.55, 11.5, False, DARK, valign=MSO_ANCHOR.MIDDLE)
add_box_text(s, "拟开展", 0.65, 5.67, 1.05, 0.43, ORANGE, None, 10.5, True, WHITE)
add_text(s, "正式 SFT、完整多粒度反馈训练、2×2 对照、3 seeds 与公开集验证",
         1.90, 5.61, 7.25, 0.55, 11.5, False, DARK, valign=MSO_ANCHOR.MIDDLE)
add_text(s, "优化器依据开发集稳定性选择；论文题目与贡献不绑定具体算法。",
         0.72, 6.55, 8.55, 0.30, 10.3, False, GRAY, PP_ALIGN.CENTER)


# 10. Completed RadarKG and retrieval
s = new_slide("已完成：RadarKG 与图文检索环境提供可执行基础", "已有研究基础", 10,
              ("已完成", TEAL))
add_metric(s, "22,241", "知识图谱边", 0.55, 1.42, 1.62, PALE_BLUE, BLUE)
add_metric(s, "10,744", "实体", 2.32, 1.42, 1.62, PALE_TEAL, TEAL)
add_metric(s, "16,609", "属性记录", 4.09, 1.42, 1.62, PALE_ORANGE, ORANGE)
add_metric(s, "1,432", "叙述文本段", 5.86, 1.42, 1.62, LIGHT, NAVY)
add_metric(s, "21,928", "策略环境过滤边", 7.63, 1.42, 1.62, PALE_RED, RED)

add_section_label(s, "已实现的混合检索链路", 0.52, 2.86, 3.1, BLUE)
pipeline = [("BM25", BLUE), ("FAISS", TEAL), ("RRF 融合", NAVY),
            ("图扩展", ORANGE), ("Cross-Encoder\n重排", RED)]
for i, (name, color) in enumerate(pipeline):
    x = 0.58 + i * 1.78
    add_box_text(s, name, x, 3.45, 1.40, 0.70, color, None, 10.5, True, WHITE)
    if i < 4:
        add_arrow(s, x + 1.45, 3.69, 0.26, 0.20, MID)

add_box_text(s, "知识图谱", 1.05, 4.72, 1.35, 0.48, PALE_BLUE, BLUE, 11, True, BLUE)
add_text(s, "结构化事实 · 多跳路径 · 集合与精确统计", 2.58, 4.68, 3.10, 0.56,
         11, False, DARK, valign=MSO_ANCHOR.MIDDLE)
add_box_text(s, "文本库", 1.05, 5.43, 1.35, 0.48, PALE_TEAL, TEAL, 11, True, TEAL)
add_text(s, "工作原理 · 内部结构 · 背景描述", 2.58, 5.39, 3.10, 0.56,
         11, False, DARK, valign=MSO_ANCHOR.MIDDLE)
add_box_text(s, "统一知识环境", 6.05, 4.81, 2.62, 0.90, NAVY, None, 14, True, WHITE)
add_arrow(s, 5.65, 5.15, 0.33, 0.22, TEAL)
add_text(s, "图检索为主干，文本检索为可选补充动作。",
         0.82, 6.48, 8.35, 0.34, 10.7, False, GRAY, PP_ALIGN.CENTER)


# 11. Completed data and executor
s = new_slide("已完成：复杂问答数据与类型化执行器具备训练基础", "已有研究基础", 11,
              ("已完成", TEAL))
add_section_label(s, "复杂问答与轨迹数据", 0.50, 1.31, 3.0, BLUE)
add_metric(s, "4,277", "Train", 0.58, 1.82, 1.35, PALE_BLUE, BLUE)
add_metric(s, "582", "Dev", 2.05, 1.82, 1.35, PALE_TEAL, TEAL)
add_metric(s, "542", "Test", 3.52, 1.82, 1.35, PALE_ORANGE, ORANGE)
add_metric(s, "4,277", "成功轨迹", 0.58, 3.05, 1.35, PALE_TEAL, TEAL)
add_metric(s, "5,401", "Oracle 回放", 2.05, 3.05, 1.35, PALE_BLUE, BLUE)
add_metric(s, "7 类", "题型覆盖", 3.52, 3.05, 1.35, LIGHT, NAVY)

add_section_label(s, "类型化执行器", 5.15, 1.31, 2.4, TEAL)
operator_rows = [["查询", "constraint · path"], ["集合", "intersect · union · difference"],
                 ["计算", "count · enumerate · aggregate"], ["判断", "compare · contains"]]
add_table(s, ["能力", "已实现节点"], operator_rows, 5.22, 1.82, 4.02, 2.38,
          [1.05, 2.97], 9.5, TEAL)

add_section_label(s, "小规模组合原型评测", 0.50, 4.63, 3.2, ORANGE)
add_metric(s, "100%", "计划有效率", 0.65, 5.08, 1.62, PALE_BLUE, BLUE)
add_metric(s, "100%", "执行成功率", 2.43, 5.08, 1.62, PALE_TEAL, TEAL)
add_metric(s, "98.6%", "答案正确率", 4.21, 5.08, 1.62, PALE_ORANGE, ORANGE,
           "71 / 72")
add_box_text(s, "边界", 6.15, 5.04, 0.72, 0.43, RED, None, 10.5, True, WHITE)
add_text(s, "72 题仅证明执行原型可行；\n不能替代正式 SFT/RL 与组合隔离实验。",
         7.02, 4.94, 2.20, 0.86, 10.2, False, RED, PP_ALIGN.CENTER,
         MSO_ANCHOR.MIDDLE)


# 12. Existing results and boundaries
s = new_slide("已有结果验证了方向，也暴露了跨域与覆盖瓶颈", "已有结果与边界", 12,
              ("非新方法结果", RED))
add_section_label(s, "RadarKG-QA-499：已有固定策略方法", 0.52, 1.30, 4.4, BLUE)
radar = [("单轮基线", 52.7, GRAY), ("RoG-style", 60.3, TEAL),
         ("固定策略路由", 88.6, BLUE)]
for i, (name, value, color) in enumerate(radar):
    y = 1.90 + i * 0.69
    add_text(s, name, 0.66, y + 0.08, 1.22, 0.28, 10.2, False, DARK)
    add_shape(s, 1.92, y, 3.10, 0.40, LIGHT, None, MSO_SHAPE.RECTANGLE)
    add_shape(s, 1.92, y, 3.10 * value / 100, 0.40, color, None, MSO_SHAPE.RECTANGLE)
    add_text(s, f"{value:.1f}%", 5.10, y + 0.02, 0.62, 0.30, 10.5, True, color)
add_text(s, "说明：类型化/结构化检索在领域组合题上有工程价值。",
         0.66, 4.15, 4.55, 0.34, 10.5, True, BLUE, PP_ALIGN.CENTER)

add_section_label(s, "KQA Pro-502：已有跨域实验", 5.35, 1.30, 3.8, ORANGE)
kqa = [("Baseline", 28.3, GRAY), ("Strategy-oracle", 28.1, ORANGE)]
for i, (name, value, color) in enumerate(kqa):
    y = 2.05 + i * 0.95
    add_text(s, name, 5.52, y + 0.09, 1.36, 0.28, 10.2, False, DARK)
    add_shape(s, 6.90, y, 1.85, 0.43, LIGHT, None, MSO_SHAPE.RECTANGLE)
    add_shape(s, 6.90, y, 1.85 * value / 40, 0.43, color, None, MSO_SHAPE.RECTANGLE)
    add_text(s, f"{value:.1f}%", 8.83, y + 0.03, 0.58, 0.30, 10.5, True, color)
add_text(s, "Δ = −0.2 pp  [−2.4, +2.0]，总体持平",
         5.54, 4.11, 3.70, 0.34, 10.5, True, ORANGE, PP_ALIGN.CENTER)

add_box_text(s, "研究启示", 0.68, 5.05, 1.10, 0.46, NAVY, None, 11, True, WHITE)
add_bullets(s, ["固定策略在领域内有效，但不能证明可学习、可迁移的组合策略",
                "跨域结果必须拆分环境覆盖、策略执行和最终答案，避免错误归因",
                "因此需要类型化策略学习、组合隔离测试与公开数据验证"],
            1.96, 4.93, 7.15, 1.42, 11.2, bullet_color=ORANGE, gap=4)
add_text(s, "注：本页均为既有系统结果，不是拟研究方法的 SFT/RL 结果。",
         0.72, 6.60, 8.55, 0.28, 9.5, True, RED, PP_ALIGN.CENTER)


# 13. Experiment design
s = new_slide("2×2 对照分别验证动作空间设计与执行反馈学习", "实验设计", 13,
              ("拟开展", ORANGE))
# Axis labels.
add_text(s, "训练方式", 0.60, 1.42, 0.85, 0.32, 10, True, GRAY, PP_ALIGN.CENTER)
add_text(s, "动作空间 →", 1.52, 1.42, 1.05, 0.32, 10, True, GRAY, PP_ALIGN.CENTER)
add_text(s, "扁平工具调用", 3.05, 1.42, 2.05, 0.32, 11.5, True, RED, PP_ALIGN.CENTER)
add_text(s, "类型化 / 分层决策", 5.28, 1.42, 2.22, 0.32, 11.5, True, TEAL, PP_ALIGN.CENTER)

cells = [
    ("SFT", "扁平 SFT", "成功轨迹模仿基线", 3.05, 1.86, PALE_RED, RED),
    ("SFT", "类型化 / 分层 SFT", "验证 action-space design", 5.28, 1.86, PALE_TEAL, TEAL),
    ("RL", "扁平 RL", "扁平空间策略优化上限", 3.05, 3.05, PALE_ORANGE, ORANGE),
    ("RL", "完整方法 RL", "类型化 + 多粒度反馈", 5.28, 3.05, PALE_BLUE, BLUE),
]
for train, name, note, x, y, fill, color in cells:
    add_box_text(s, name, x, y, 2.02, 0.58, color, None, 10.8, True, WHITE)
    add_box_text(s, note, x, y + 0.61, 2.02, 0.49, fill, color, 9.2, False, color)
add_text(s, "SFT", 1.90, 2.12, 0.72, 0.32, 11.5, True, GRAY, PP_ALIGN.CENTER)
add_text(s, "SFT + RL", 1.73, 3.34, 1.05, 0.32, 11.5, True, GRAY, PP_ALIGN.CENTER)
add_arrow(s, 5.08, 2.25, 0.16, 0.18, TEAL)
add_arrow(s, 5.08, 3.43, 0.16, 0.18, BLUE)
add_arrow(s, 4.00, 2.90, 0.20, 0.18, ORANGE, "down")
add_arrow(s, 6.23, 2.90, 0.20, 0.18, BLUE, "down")

add_section_label(s, "四类测试", 0.52, 4.48, 1.8, NAVY)
tests = [("IID", "总体效果"), ("实体隔离", "实体泛化"),
         ("组合隔离", "策略 ≠ 模板",), ("错误恢复", "利用环境反馈")]
for i, (a, b) in enumerate(tests):
    x = 0.58 + i * 1.65
    fill = PALE_ORANGE if a == "组合隔离" else WHITE
    line = ORANGE if a == "组合隔离" else LIGHT
    add_box_text(s, a + "\n" + b, x, 4.96, 1.46, 0.72, fill, line, 9.7,
                 a == "组合隔离", ORANGE if a == "组合隔离" else DARK)

add_section_label(s, "公开验证", 7.22, 4.48, 1.8, TEAL)
add_box_text(s, "KQA Pro\n类型化程序 · 集合 · 计数", 7.28, 4.96, 1.05, 0.72,
             PALE_TEAL, TEAL, 8.8, True, TEAL)
add_box_text(s, "MetaQA\n1 / 2 / 3-hop 路径", 8.44, 4.96, 1.02, 0.72,
             PALE_BLUE, BLUE, 8.8, True, BLUE)
add_text(s, "分层报告：环境 oracle  →  策略执行  →  最终答案",
         0.75, 6.15, 8.55, 0.35, 11.5, True, NAVY, PP_ALIGN.CENTER)
add_text(s, "指标：EM / 集合 F1 / 计划有效率 / 支持事实覆盖 / 重规划成功率 / 调用数与延迟；至少 3 seeds。",
         0.72, 6.57, 8.60, 0.27, 9.2, False, GRAY, PP_ALIGN.CENTER)


# 14. Innovations, feasibility, risks
s = new_slide("两项创新可独立验证，主要风险已有兜底方案", "创新与可行性", 14)
add_section_label(s, "两项核心方法创新", 0.50, 1.28, 3.0, BLUE)
add_box_text(s, "创新 1", 0.60, 1.80, 0.90, 0.42, BLUE, None, 10.5, True, WHITE)
add_text(s, "复杂问答的类型化图检索决策建模", 1.70, 1.76, 3.08, 0.50,
         13.2, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
add_text(s, "以宏动作、结构化参数、中间变量和确定性执行解决 action-space design。",
         0.72, 2.42, 4.00, 0.72, 10.8, False, DARK)
add_box_text(s, "创新 2", 0.60, 3.31, 0.90, 0.42, ORANGE, None, 10.5, True, WHITE)
add_text(s, "基于多粒度可执行反馈的策略学习", 1.70, 3.27, 3.08, 0.50,
         13.2, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
add_text(s, "利用计划、证据、进展、恢复与成本反馈解决长程 credit assignment。",
         0.72, 3.93, 4.00, 0.72, 10.8, False, DARK)
add_box_text(s, "组合泛化是验证手段，不包装为第三项算法创新", 0.70, 5.10, 3.95, 0.55,
             PALE_TEAL, TEAL, 10.5, True, TEAL)

add_section_label(s, "风险与对策", 5.05, 1.28, 2.3, RED)
risks = [
    ["RL 相对 SFT 无增益", "强化组合隔离/恢复任务；检查反馈方差与门控"],
    ["过程奖励被投机", "终局答案门控；重复惩罚；轨迹审计"],
    ["KG 缺失造成错误负奖", "可闭世界子集训练；报告环境 oracle 上限"],
    ["公开集迁移无增益", "按题型与覆盖分层分析，不选择性报告"],
    ["显存与训练成本不足", "1.5B/3B LoRA 验证；量化与 offload"]
]
add_table(s, ["风险", "对策"], risks, 5.10, 1.78, 4.25, 3.68,
          [1.55, 2.70], 8.4, RED)
add_box_text(s, "可行性", 5.18, 5.80, 0.82, 0.42, TEAL, None, 10.5, True, WHITE)
add_text(s, "知识环境、执行器、数据、成功轨迹与训练脚本原型均已具备",
         6.18, 5.74, 3.02, 0.56, 10.5, True, TEAL, PP_ALIGN.CENTER,
         MSO_ANCHOR.MIDDLE)
add_text(s, "来源/置信度仅作可追溯展示与可选扩展，不承担论文成立条件。",
         5.16, 6.52, 4.16, 0.30, 9.2, False, GRAY, PP_ALIGN.CENTER)


# 15. Schedule and deliverables
s = new_slide("研究计划围绕“建模—学习—验证—写作”形成闭环", "后续计划", 15,
              ("10 个月", BLUE))
add_section_label(s, "研究进度", 0.48, 1.25, 1.8, BLUE)
months = list(range(1, 11))
left = 2.45; cell_w = 0.54
for j, m in enumerate(months):
    add_box_text(s, str(m), left + j * cell_w, 1.62, cell_w - 0.02, 0.38,
                 NAVY, None, 8.5, True, WHITE, shape=MSO_SHAPE.RECTANGLE)
phases = [
    ("问题与数据冻结", 1, 2, BLUE),
    ("类型化环境与 2×2 基线", 3, 4, TEAL),
    ("SFT / 策略优化 / 消融", 4, 6, ORANGE),
    ("雷达与公开集实验", 7, 8, BLUE),
    ("系统整合与论文写作", 9, 10, TEAL),
]
for i, (name, start, end, color) in enumerate(phases):
    y = 2.15 + i * 0.56
    add_text(s, name, 0.57, y + 0.02, 1.72, 0.30, 9.5, False, DARK)
    for m in months:
        fill = color if start <= m <= end else LIGHT
        add_shape(s, left + (m - 1) * cell_w, y, cell_w - 0.02, 0.34,
                  fill, None, MSO_SHAPE.RECTANGLE)

add_section_label(s, "预期成果", 0.48, 5.30, 1.9, TEAL)
outputs = [("方法", "类型化图检索决策"), ("方法", "多粒度反馈策略学习"),
           ("数据", "组合隔离雷达问答集"), ("实验", "RadarKG / KQA Pro / MetaQA"),
           ("系统", "可执行轨迹与答案展示原型")]
for i, (tag, text) in enumerate(outputs):
    x = 0.58 + i * 1.82
    color = [BLUE, ORANGE, TEAL, NAVY, RED][i]
    add_box_text(s, tag, x, 5.83, 0.52, 0.42, color, None, 9.2, True, WHITE)
    add_text(s, text, x + 0.59, 5.76, 1.06, 0.56, 9.0, True, DARK,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

add_box_text(s, "主线", 0.62, 6.53, 0.72, 0.40, NAVY, None, 10, True, WHITE)
add_text(s, "类型化决策解决“如何表示”  ·  可执行反馈解决“如何学习”  ·  组合隔离验证“是否泛化”",
         1.55, 6.47, 7.70, 0.50, 11.5, True, NAVY, PP_ALIGN.CENTER,
         MSO_ANCHOR.MIDDLE)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUTPUT))
print(f"saved: {OUTPUT}")
print(f"slides: {len(prs.slides)}; size: {SW / 914400:.2f} x {SH / 914400:.2f} in")
