# -*- coding: utf-8 -*-
"""Apply terminology reduction and Chinese-first narration to the defense deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "reports" / "硕士学位论文开题答辩_图检索策略学习_定稿优化版.pptx"
OUTPUT = ROOT / "reports" / "硕士学位论文开题答辩_图检索策略学习_最终答辩版.pptx"

NAVY = RGBColor(0x18, 0x36, 0x57)
BLUE = RGBColor(0x2F, 0x64, 0x91)
GRAY = RGBColor(0x68, 0x73, 0x7D)
LINE = RGBColor(0xD8, 0xE0, 0xE7)
PALE_BLUE = RGBColor(0xEC, 0xF1, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


prs = Presentation(str(SOURCE))
prs.core_properties.subject = "硕士学位论文开题答辩最终答辩版"


def text_frames(slide):
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            yield shape.text_frame
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def replace_all(slide, replacements):
    for tf in text_frames(slide):
        for paragraph in tf.paragraphs:
            for run in paragraph.runs:
                for old, new in replacements:
                    if old in run.text:
                        run.text = run.text.replace(old, new)


def remove_text_shape(slide, exact_text):
    for shape in list(slide.shapes):
        if getattr(shape, "has_text_frame", False) and shape.text.strip() == exact_text:
            slide.shapes._spTree.remove(shape._element)


def find_shape(slide, exact_text):
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip() == exact_text:
            return shape
    return None


def style_shape_text(shape, size, color, bold):
    if shape is None:
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.bold = bold


# Slide 4: Chinese-first Method Overview -----------------------------------
s = prs.slides[3]
remove_text_shape(s, "a_t ~ πθ(a_t | s_t)")
remove_text_shape(s, "s_t = {q, history, observation}")
replace_all(s, [
    ("METHOD OVERVIEW  |  总体研究思路", "总体研究思路"),
    ("确定性执行环境\nGraph Retrieval · Set\nCount / Compare · Text",
     "确定性执行环境\n图检索 · 集合运算\n计数/比较 · 文本检索"),
    ("SFT\nInitialization", "SFT\n冷启动"),
    ("Initial Policy\nπθ0", "初始策略\nπθ0"),
    ("Retrieval Policy\nπθ", "检索策略\nπθ"),
    ("Rollout\nTrajectory τ", "交互轨迹\nτ"),
    ("Executable\nFeedback R(τ)", "执行反馈\nR(τ)"),
    ("Policy\nOptimization", "策略优化"),
    ("update πθ", "更新 πθ"),
    ("Answer · Plan · Progress · Recovery · Cost", "答案 · 计划 · 进展 · 恢复 · 成本"),
    ("Generalization\nEvaluation", "泛化验证"),
    ("Composition Holdout\nKQA Pro\nMetaQA", "组合隔离\nKQA Pro\nMetaQA"),
    ("ZOOM-IN", "方法展开"),
    ("Policy Representation\n第 5 页 · 策略如何表示", "第 5 页 · 策略表示\n“策略如何表示？”"),
    ("Policy Learning\n第 6 页 · 策略如何学习", "第 6 页 · 策略学习\n“策略如何学习？”"),
])


# Slide 5: emphasize Entity -> EntitySet -> Scalar --------------------------
s = prs.slides[4]
replace_all(s, [
    ("macro: 路径 / 约束 / 集合 / 聚合    operator: constraint / path / intersect / count",
     "macro：任务类型（路径 / 约束 / 集合 / 聚合）    operator：执行算子（constraint / path / intersect / count）"),
    ("args: 实体 / 关系 / 属性    bindings: 中间变量    control: continue / revise / finish",
     "args：实体 / 关系 / 属性    bindings：中间变量    control：继续 / 修正 / 结束"),
])
for label in ("Entity\n实体", "EntitySet\n实体集合", "Scalar\n标量"):
    shape = find_shape(s, label)
    if shape is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = PALE_BLUE
        shape.line.color.rgb = BLUE
        shape.line.width = Pt(1.1)
        style_shape_text(shape, 9.0, NAVY, True)
for label in ("Relation\n关系", "Evidence\n证据"):
    shape = find_shape(s, label)
    if shape is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = LINE
        shape.line.width = Pt(0.6)
        style_shape_text(shape, 7.7, GRAY, False)


# Slide 6: Chinese-first Policy Learning ------------------------------------
s = prs.slides[5]
replace_all(s, [
    ("POLICY LEARNING  |  策略优化", "策略学习与优化"),
    ("SFT Initialization", "SFT 初始化"),
    ("SFT\nInitialization", "SFT\n冷启动"),
    ("Initial Policy\nπθ0", "初始策略\nπθ0"),
    ("Policy Learning Loop", "策略学习闭环"),
    ("Retrieval Policy\nπθ", "检索策略\nπθ"),
    ("Executable\nEnvironment", "可执行环境"),
    ("Trajectory\nτ", "完整轨迹\nτ"),
    ("Executable\nEvaluation", "可执行评价"),
    ("Reward\nR(τ)", "奖励\nR(τ)"),
    ("Policy\nOptimization", "策略优化"),
    ("update πθ", "更新 πθ"),
    ("Reward Decomposition", "奖励分解"),
    ("Answer · Plan · Progress\nRecovery · Cost", "答案 · 计划 · 进展\n恢复 · 成本"),
    ("TRAINING CURRICULUM", "课程式训练"),
    ("执行环境产生轨迹，轨迹形成可执行反馈，反馈持续更新同一个检索策略 πθ。",
     "执行环境产生轨迹，轨迹形成可计算反馈，并持续更新同一个检索策略 πθ。"),
])


# Slide 9: Chinese 2x2 experiment terminology -------------------------------
s = prs.slides[8]
replace_all(s, [
    ("Flat Policy Space", "扁平策略空间"),
    ("Typed Policy Space", "类型化策略空间"),
    ("Flat-SFT", "扁平 SFT"),
    ("Typed-SFT", "类型化 SFT"),
    ("Flat-RL", "扁平 RL"),
    ("Full Method", "完整方法"),
    ("IID / Entity Holdout / Composition Holdout / Error Recovery",
     "IID / 实体隔离 / 组合隔离 / 错误恢复"),
    ("Environment Oracle  →  Policy Execution  →  Final Answer  →  Cost / Recovery",
     "环境上限（Oracle） → 策略执行 → 最终答案 → 成本 / 恢复"),
    ("横向验证策略表示，纵向验证策略优化，组合隔离检验 compositional generalization。",
     "横向验证策略表示，纵向验证策略优化，组合隔离检验组合泛化能力。"),
])
shape = find_shape(s, "Composition Holdout")
if shape is not None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace("Composition Holdout", "组合隔离\n（Composition Holdout）")
            run.font.size = Pt(7.5)
            run.font.color.rgb = BLUE
            run.font.bold = True
    shape.top = Inches(4.29)
    shape.height = Inches(0.38)
definition = find_shape(s, "训练与测试采用不同算子组合 / 程序结构，检验可组合策略而非模板记忆。")
if definition is not None:
    definition.top = Inches(4.72)


# Slide 10: Chinese causal-comparison names ---------------------------------
s = prs.slides[9]
replace_all(s, [
    ("Flat-SFT → Typed-SFT", "扁平 SFT → 类型化 SFT"),
    ("Typed-SFT → Full-RL", "类型化 SFT → 完整方法 RL"),
])


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUTPUT))
print(f"saved: {OUTPUT}")
print(f"slides: {len(prs.slides)}; language-polished: 4, 5, 6, 9, 10")
