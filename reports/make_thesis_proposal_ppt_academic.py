# -*- coding: utf-8 -*-
"""Generate the academically restyled thesis proposal deck.

Structure: 13 main slides + 2 backup slides. All diagrams remain editable.
"""

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "PPT模板1.pptx"
OUTPUT = ROOT / "reports" / "硕士学位论文开题答辩_图检索策略学习_学术重构版.pptx"

# Restrained academic palette: one school/navy blue, one blue-gray auxiliary,
# and a muted gold used only for the most important comparison or conclusion.
NAVY = RGBColor(0x1B, 0x36, 0x55)
SCHOOL_BLUE = RGBColor(0x2E, 0x5F, 0x8A)
BLUE_GRAY = RGBColor(0x6F, 0x88, 0x9E)
AUX_TEAL = RGBColor(0x5F, 0x87, 0x84)
GOLD = RGBColor(0xB5, 0x8B, 0x3C)
INK = RGBColor(0x23, 0x2C, 0x35)
GRAY = RGBColor(0x67, 0x72, 0x7E)
MID = RGBColor(0xB9, 0xC4, 0xCE)
LINE = RGBColor(0xD7, 0xDF, 0xE6)
PALE_BLUE = RGBColor(0xEB, 0xF0, 0xF4)
PALE_TEAL = RGBColor(0xEC, 0xF2, 0xF1)
PALE_GOLD = RGBColor(0xF6, 0xF1, 0xE7)
BG = RGBColor(0xF8, 0xF9, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "微软雅黑"
MONO = "Consolas"


prs = Presentation(str(TEMPLATE))
SW, SH = prs.slide_width, prs.slide_height
logo_shape = next(sh for sh in prs.slides[0].shapes if sh.shape_type == 13)
LOGO = BytesIO(logo_shape.image.blob)
while len(prs.slides):
    sid = prs.slides._sldIdLst[0]
    prs.part.drop_rel(sid.rId)
    del prs.slides._sldIdLst[0]
BLANK = prs.slide_layouts[6]

prs.core_properties.title = "面向复杂雷达情报问答的图检索策略学习方法研究"
prs.core_properties.subject = "硕士学位论文开题答辩（学术重构版）"
prs.core_properties.author = ""


def set_typeface(run, name=FONT):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", name)


def run_style(run, size=13, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_typeface(run, name)


def text(slide, value, x, y, w, h, size=13, bold=False, color=INK,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, name=FONT,
         margin=0.03):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value; run_style(r, size, bold, color, name)
    return box


def bullets(slide, items, x, y, w, h, size=12, color=INK, gap=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        value, level = item if isinstance(item, tuple) else (item, 0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.left_margin = Inches(0.16 + 0.20 * level)
        p.first_line_indent = Inches(-0.13)
        rb = p.add_run(); rb.text = ("•" if level == 0 else "–") + " "
        run_style(rb, size - level, True, SCHOOL_BLUE if level == 0 else BLUE_GRAY)
        rv = p.add_run(); rv.text = value
        run_style(rv, size - level, False, color if level == 0 else GRAY)
    return box


def shape(slide, x, y, w, h, fill=WHITE, line=LINE,
          kind=MSO_SHAPE.RECTANGLE, width=0.9):
    obj = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid(); obj.fill.fore_color.rgb = fill
    if line is None:
        obj.line.fill.background()
    else:
        obj.line.color.rgb = line; obj.line.width = Pt(width)
    obj.shadow.inherit = False
    return obj


def box(slide, value, x, y, w, h, fill=WHITE, line=LINE, size=12,
        bold=False, color=INK, align=PP_ALIGN.CENTER,
        kind=MSO_SHAPE.RECTANGLE, line_width=0.9, margin=0.06):
    obj = shape(slide, x, y, w, h, fill, line, kind, line_width)
    tf = obj.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value; run_style(r, size, bold, color)
    return obj


def line(slide, x1, y1, x2, y2, color=MID, width=1.3):
    obj = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                     Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    obj.line.color.rgb = color; obj.line.width = Pt(width)
    return obj


def arrow(slide, x, y, w, h, color=BLUE_GRAY, direction="right"):
    kind = {"right": MSO_SHAPE.RIGHT_ARROW, "left": MSO_SHAPE.LEFT_ARROW,
            "down": MSO_SHAPE.DOWN_ARROW, "up": MSO_SHAPE.UP_ARROW}[direction]
    return shape(slide, x, y, w, h, color, None, kind)


def title(slide, main, section, page, appendix=False):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    shape(slide, 0, 0, 0.08, 7.5, NAVY, None)
    text(slide, section, 0.46, 0.18, 2.85, 0.25, 9.2, True, AUX_TEAL)
    text(slide, main, 0.46, 0.48, 7.25, 0.58, 21, True, NAVY)
    line(slide, 0.46, 1.10, 9.57, 1.10, MID, 0.9)
    LOGO.seek(0); slide.shapes.add_picture(LOGO, Inches(7.70), Inches(0.12), width=Inches(2.12))
    line(slide, 0.46, 7.15, 9.57, 7.15, LINE, 0.7)
    text(slide, "备份材料" if appendix else "硕士学位论文开题答辩  ·  图检索策略学习",
         0.48, 7.19, 6.3, 0.19, 8.3, False, GRAY)
    text(slide, f"{page:02d}", 9.12, 7.18, 0.36, 0.20, 8.5, True, GRAY, PP_ALIGN.RIGHT)
    return slide


def new_slide(main, section, page, appendix=False):
    return title(prs.slides.add_slide(BLANK), main, section, page, appendix)


def section_label(slide, value, x, y, w, color=SCHOOL_BLUE):
    shape(slide, x, y + 0.03, 0.045, 0.28, color, None)
    text(slide, value, x + 0.13, y, w - 0.13, 0.34, 11.2, True, color)


def metric(slide, value, label, x, y, w=1.55, accent=SCHOOL_BLUE):
    shape(slide, x, y, w, 0.93, WHITE, LINE)
    text(slide, value, x + 0.05, y + 0.10, w - 0.10, 0.38, 20, True, accent, PP_ALIGN.CENTER)
    text(slide, label, x + 0.05, y + 0.55, w - 0.10, 0.23, 9.7, False, GRAY, PP_ALIGN.CENTER)


def table(slide, headers, rows, x, y, w, h, widths=None, size=9.6,
          header=NAVY, highlight=()):
    t = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y),
                               Inches(w), Inches(h)).table
    if widths:
        for i, v in enumerate(widths): t.columns[i].width = Inches(v)
    for j, value in enumerate(headers):
        c = t.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = header
        c.margin_left = c.margin_right = Inches(0.03)
        c.margin_top = c.margin_bottom = Inches(0.02); c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = value; run_style(r, size, True, WHITE)
    for i, row in enumerate(rows, 1):
        for j, value in enumerate(row):
            c = t.cell(i, j); c.fill.solid()
            c.fill.fore_color.rgb = PALE_GOLD if i in highlight else (WHITE if i % 2 else PALE_BLUE)
            c.margin_left = c.margin_right = Inches(0.03)
            c.margin_top = c.margin_bottom = Inches(0.02); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run(); r.text = str(value)
            run_style(r, size, i in highlight, GOLD if i in highlight else INK)
    return t


# 1 Cover -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
shape(s, 0, 0, 0.12, 7.5, NAVY, None)
LOGO.seek(0); s.shapes.add_picture(LOGO, Inches(6.95), Inches(0.12), width=Inches(2.85))
text(s, "硕士学位论文开题答辩", 0.78, 1.45, 2.50, 0.34, 11, True, AUX_TEAL)
text(s, "面向复杂雷达情报问答的\n图检索策略学习方法研究",
     0.76, 2.05, 8.45, 1.38, 29, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
text(s, "Graph Retrieval Policy Learning for Complex Radar Intelligence Question Answering",
     0.80, 3.64, 8.30, 0.40, 12, False, GRAY)
shape(s, 0.80, 4.32, 3.05, 0.035, SCHOOL_BLUE, None)
text(s, "汇报人：XXX    指导教师：XXX\n专业：XXX      日期：2026 年  月  日",
     0.80, 4.68, 5.70, 0.85, 14.5, False, INK)
box(s, "研究问题", 6.95, 5.10, 0.92, 0.38, PALE_BLUE, SCHOOL_BLUE, 9.5, True, SCHOOL_BLUE)
text(s, "复杂组合检索的决策表示与策略学习", 7.98, 5.05, 1.50, 0.55,
     10.5, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, "类型化可执行决策  ·  多粒度执行反馈  ·  组合泛化验证",
     0.80, 6.62, 8.30, 0.32, 12.5, True, SCHOOL_BLUE)


# 2 Background ---------------------------------------------------------------
s = new_slide("复杂雷达问答的核心需求是组合检索", "研究背景", 2)
section_label(s, "异构知识环境", 0.52, 1.35, 2.0)
for i, (a, b) in enumerate([("装备手册", "PDF / 扫描文档"), ("知识图谱", "实体 · 关系 · 属性"),
                             ("叙述文本", "原理 · 结构 · 背景")]):
    y = 1.83 + i * 0.94
    box(s, a, 0.60, y, 1.25, 0.55, PALE_BLUE, SCHOOL_BLUE, 11, True, NAVY)
    text(s, b, 2.02, y + 0.04, 1.48, 0.44, 10, False, GRAY, valign=MSO_ANCHOR.MIDDLE)
    arrow(s, 3.48, y + 0.18, 0.28, 0.16)

box(s, "复杂问题", 3.88, 2.44, 1.18, 0.70, NAVY, None, 13, True, WHITE)
text(s, "“找出美国研制、工作在 J 波段的雷达，\n并统计数量与最大探测距离”",
     3.42, 3.37, 2.08, 0.84, 10.8, False, INK, PP_ALIGN.CENTER)

section_label(s, "组合查询结构", 5.75, 1.35, 2.1)
items = [("多跳路径", "Radar → Platform → Country"), ("集合运算", "交 ∩ / 并 ∪ / 差 −"),
         ("精确计算", "Count / Aggregate / Compare"), ("图文协同", "结构化事实 + 解释性文本")]
for i, (a, b) in enumerate(items):
    y = 1.80 + i * 0.87
    text(s, f"0{i+1}", 5.83, y + 0.08, 0.36, 0.28, 9.5, True, AUX_TEAL, PP_ALIGN.CENTER)
    text(s, a, 6.28, y, 1.15, 0.40, 11.5, True, NAVY)
    text(s, b, 7.50, y, 1.82, 0.45, 9.8, False, GRAY)
    line(s, 6.24, y + 0.49, 9.20, y + 0.49, LINE, 0.7)

shape(s, 0.66, 5.70, 8.68, 0.86, WHITE, LINE)
text(s, "研究对象", 0.88, 5.91, 1.05, 0.30, 10.5, True, AUX_TEAL)
text(s, "自然语言复杂问题如何转化为可执行、可反馈、可泛化的图检索策略",
     2.02, 5.80, 7.02, 0.49, 15, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


# 3 Challenges ---------------------------------------------------------------
s = new_slide("复杂图检索面临动作空间与学习信号两个核心挑战", "问题提出", 3)
box(s, "挑战一", 0.62, 1.48, 0.88, 0.42, SCHOOL_BLUE, None, 10.5, True, WHITE)
text(s, "Action-space design", 1.67, 1.46, 2.50, 0.40, 12, True, BLUE_GRAY)
text(s, "扁平工具空间将多种决策耦合在同一步骤", 0.68, 2.10, 3.82, 0.48, 15, True, NAVY)
bullets(s, ["任务分解、工具选择、参数生成与精确计算同时由 LLM 完成",
            "工具、参数与历史状态组合增长，长程调用更难保持一致",
            "需要结构化动作与确定性执行分工"], 0.72, 2.78, 3.72, 1.75, 11.5)

arrow(s, 4.58, 3.12, 0.58, 0.30, GOLD)

box(s, "挑战二", 5.28, 1.48, 0.88, 0.42, SCHOOL_BLUE, None, 10.5, True, WHITE)
text(s, "Credit assignment", 6.33, 1.46, 2.50, 0.40, 12, True, BLUE_GRAY)
text(s, "最终答案奖励难以定位中间决策贡献", 5.33, 2.10, 3.88, 0.48, 15, True, NAVY)
bullets(s, ["Answer EM/F1 只评价整条轨迹，复杂任务中奖励稀疏",
            "计划、证据、状态进展与错误修复均可由环境观测",
            "需要将可执行过程转化为细粒度学习信号"], 5.38, 2.78, 3.70, 1.75, 11.5)

line(s, 0.72, 5.10, 9.20, 5.10, LINE, 0.8)
text(s, "现有 RAG / GraphRAG / 工具 Agent 分别增强了知识召回、图结构利用和环境交互；",
     0.85, 5.43, 8.25, 0.34, 11, False, GRAY, PP_ALIGN.CENTER)
text(s, "本文进一步研究复杂组合检索的决策表示与策略学习。",
     0.85, 5.91, 8.25, 0.40, 13.5, True, SCHOOL_BLUE, PP_ALIGN.CENTER)
text(s, "参考：RAG (NeurIPS 2020)；ReAct (ICLR 2023)；ToG / RoG (ICLR 2024)；Search-R1 / Graph-R1。",
     0.72, 6.62, 8.65, 0.26, 8.3, False, GRAY, PP_ALIGN.CENTER)


# 4 Questions ----------------------------------------------------------------
s = new_slide("两个科学问题形成“决策表示 → 策略学习”的递进关系", "科学问题", 4)
box(s, "Q1", 0.72, 1.48, 0.58, 0.58, SCHOOL_BLUE, None, 13, True, WHITE, kind=MSO_SHAPE.OVAL)
text(s, "复杂图检索决策如何被可执行地表示？", 1.52, 1.46, 3.45, 0.56,
     16, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
shape(s, 0.78, 2.35, 4.02, 2.35, WHITE, LINE)
text(s, "方法一", 1.00, 2.62, 0.80, 0.30, 10, True, AUX_TEAL)
text(s, "类型化图检索决策建模", 1.88, 2.55, 2.60, 0.45, 14, True, SCHOOL_BLUE)
bullets(s, ["语义宏动作与结构化参数", "中间变量和数据依赖", "确定性集合、路径与聚合执行"],
        1.02, 3.22, 3.54, 1.15, 11.2)

arrow(s, 4.90, 3.35, 0.40, 0.24, GOLD)

box(s, "Q2", 5.40, 1.48, 0.58, 0.58, SCHOOL_BLUE, None, 13, True, WHITE, kind=MSO_SHAPE.OVAL)
text(s, "长程检索策略如何获得有效学习信号？", 6.20, 1.46, 3.18, 0.56,
     16, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
shape(s, 5.46, 2.35, 3.92, 2.35, WHITE, LINE)
text(s, "方法二", 5.68, 2.62, 0.80, 0.30, 10, True, AUX_TEAL)
text(s, "多粒度可执行反馈策略学习", 6.55, 2.55, 2.55, 0.45, 14, True, SCHOOL_BLUE)
bullets(s, ["计划合法性与支持事实", "状态进展和错误恢复", "答案质量与执行成本联合优化"],
        5.70, 3.22, 3.42, 1.15, 11.2)

shape(s, 1.08, 5.42, 7.82, 0.78, PALE_GOLD, GOLD)
text(s, "总体目标：学习可执行、经济且具有组合泛化能力的图检索策略",
     1.25, 5.56, 7.47, 0.44, 15, True, NAVY, PP_ALIGN.CENTER,
     MSO_ANCHOR.MIDDLE)


# 5 Overall framework ---------------------------------------------------------
s = new_slide("总体框架连接推理执行与策略学习两个闭环", "总体技术路线", 5)
section_label(s, "推理执行闭环", 0.50, 1.30, 2.0)
stages = [("自然语言\n复杂问题", 0.48, 1.05, NAVY),
          ("语义 / 策略\n决策", 1.74, 1.10, SCHOOL_BLUE),
          ("类型化\n查询计划", 3.07, 1.06, SCHOOL_BLUE),
          ("确定性执行\n图工具 · 文本工具", 4.36, 1.34, NAVY),
          ("结构化观察\n结果 · 状态 · 错误", 5.92, 1.30, BLUE_GRAY)]
for value, x, w, color in stages:
    box(s, value, x, 1.78, w, 0.90, color, None, 10.7, True, WHITE)
for x in [1.56, 2.87, 4.17, 5.72]: arrow(s, x, 2.11, 0.18, 0.18)
arrow(s, 7.29, 2.10, 0.22, 0.20, GOLD)
box(s, "最终答案", 7.56, 1.82, 1.28, 0.82, PALE_GOLD, GOLD, 12, True, NAVY)

# Observation loop.
line(s, 6.58, 2.75, 6.58, 3.22, AUX_TEAL, 1.5)
line(s, 6.58, 3.22, 2.30, 3.22, AUX_TEAL, 1.5)
arrow(s, 1.94, 3.09, 0.34, 0.23, AUX_TEAL, "left")
text(s, "继续规划 · 参数修正 · 终止判断", 3.38, 2.94, 2.45, 0.25,
     9.5, True, AUX_TEAL, PP_ALIGN.CENTER)

section_label(s, "策略学习闭环", 0.50, 3.73, 2.0, SCHOOL_BLUE)
box(s, "成功轨迹", 0.63, 4.25, 1.42, 0.58, PALE_BLUE, SCHOOL_BLUE, 10.5, True, NAVY)
arrow(s, 2.12, 4.44, 0.28, 0.18)
box(s, "SFT 冷启动", 2.48, 4.25, 1.42, 0.58, PALE_BLUE, SCHOOL_BLUE, 10.5, True, NAVY)
arrow(s, 3.97, 4.44, 0.28, 0.18)
box(s, "策略优化", 4.33, 4.25, 1.42, 0.58, PALE_GOLD, GOLD, 10.5, True, NAVY)
arrow(s, 5.83, 4.44, 0.28, 0.18, GOLD)
# Draw the environment-to-feedback path first so the policy box remains legible.
line(s, 6.58, 3.25, 6.58, 5.00, BLUE_GRAY, 1.2)
arrow(s, 6.47, 5.00, 0.22, 0.27, BLUE_GRAY, "down")
box(s, "检索策略 πθ", 6.08, 4.20, 1.68, 0.68, NAVY, None, 10.8, True, WHITE)

text(s, "执行反馈", 0.66, 5.31, 0.88, 0.30, 10.5, True, BLUE_GRAY)
for i, value in enumerate(["Answer", "Plan", "Progress", "Recovery", "Cost"]):
    x = 1.62 + i * 1.22
    fill = PALE_GOLD if value in ("Progress", "Recovery") else WHITE
    border = GOLD if value in ("Progress", "Recovery") else LINE
    box(s, value, x, 5.22, 1.05, 0.44, fill, border, 9.5, True,
        GOLD if value in ("Progress", "Recovery") else NAVY)
line(s, 7.42, 5.16, 7.42, 4.90, GOLD, 1.4)
arrow(s, 7.30, 4.80, 0.24, 0.27, GOLD, "up")

text(s, "策略模型负责语义决策，执行器负责确定性计算；环境观察同时支撑在线修正与离线学习。",
     0.72, 6.35, 8.58, 0.38, 10.5, False, GRAY, PP_ALIGN.CENTER)


# 6 Method one comparison -----------------------------------------------------
s = new_slide("类型化决策将语义规划与精确执行职责解耦", "方法一  ·  Action-space design", 6)
text(s, "扁平 Agent", 0.78, 1.38, 3.72, 0.38, 13, True, BLUE_GRAY, PP_ALIGN.CENTER)
shape(s, 0.62, 1.82, 4.02, 4.25, WHITE, LINE)
text(s, "单步决策空间", 0.86, 2.07, 1.25, 0.30, 10, True, GRAY)
tools = ["graph_lookup", "subgraph", "set_op", "count", "attr_filter", "text_search"]
for i, value in enumerate(tools):
    x = 0.88 + (i % 3) * 1.10; y = 2.55 + (i // 3) * 0.58
    box(s, value, x, y, 0.98, 0.38, PALE_BLUE, LINE, 8.0, False, NAVY)
text(s, "×  自由参数  ×  历史状态", 1.12, 3.91, 2.98, 0.36, 12, True, BLUE_GRAY, PP_ALIGN.CENTER)
line(s, 2.62, 4.42, 2.62, 4.80, BLUE_GRAY, 1.4)
arrow(s, 2.50, 4.74, 0.24, 0.27, BLUE_GRAY, "down")
box(s, "LLM 同时完成决策与结果计算", 1.25, 5.06, 2.74, 0.52,
    PALE_BLUE, BLUE_GRAY, 10.3, True, NAVY)

text(s, "类型化图检索决策", 5.20, 1.38, 3.86, 0.38, 13, True, SCHOOL_BLUE, PP_ALIGN.CENTER)
shape(s, 4.92, 1.82, 4.42, 4.25, WHITE, SCHOOL_BLUE, width=1.2)
steps = [("1", "语义宏动作", "路径 / 约束 / 集合 / 聚合"),
         ("2", "结构化参数", "Entity · Relation · EntitySet · Scalar"),
         ("3", "确定性执行", "交并差 · 路径 · 计数 · 比较")]
for i, (num, a, b) in enumerate(steps):
    y = 2.18 + i * 0.93
    box(s, num, 5.20, y, 0.38, 0.38, SCHOOL_BLUE, None, 9.5, True, WHITE,
        kind=MSO_SHAPE.OVAL)
    text(s, a, 5.78, y - 0.02, 1.30, 0.35, 11.2, True, NAVY)
    text(s, b, 7.08, y - 0.02, 1.92, 0.42, 9.4, False, GRAY)
    if i < 2:
        line(s, 5.39, y + 0.42, 5.39, y + 0.80, SCHOOL_BLUE, 1.2)
box(s, "结构化观察返回策略", 5.77, 5.20, 2.78, 0.48,
    PALE_GOLD, GOLD, 10.5, True, NAVY)

text(s, "动作因子化", 4.35, 6.34, 1.28, 0.32, 10.5, True, GOLD, PP_ALIGN.CENTER)
line(s, 0.90, 6.50, 4.17, 6.50, LINE, 0.8)
line(s, 5.78, 6.50, 9.10, 6.50, LINE, 0.8)
text(s, "宏动作缩小选择范围，类型系统约束参数，执行器保证集合与数值计算的一致性。",
     1.05, 6.69, 8.00, 0.28, 10.2, False, GRAY, PP_ALIGN.CENTER)


# 7 Plan example --------------------------------------------------------------
s = new_slide("类型化计划把组合查询编译为可执行数据流", "方法一  ·  可执行表示", 7)
text(s, "示例问题", 0.62, 1.35, 0.90, 0.28, 10, True, AUX_TEAL)
text(s, "“美国研制且工作在 J 波段的雷达有多少款？”",
     1.55, 1.30, 7.65, 0.42, 15, True, NAVY, PP_ALIGN.CENTER)

box(s, "count\nEntitySet → Scalar", 3.72, 2.02, 2.12, 0.68, NAVY, None, 11, True, WHITE)
box(s, "intersect\nEntitySet × EntitySet → EntitySet", 3.30, 3.10, 2.96, 0.76,
    SCHOOL_BLUE, None, 10.2, True, WHITE)
box(s, "constraint\ncountryOfOrigin = 美国", 1.12, 4.48, 2.86, 0.78,
    PALE_BLUE, SCHOOL_BLUE, 10.2, True, NAVY)
box(s, "constraint\nhasFrequencyBand = J", 5.58, 4.48, 2.86, 0.78,
    PALE_BLUE, SCHOOL_BLUE, 10.2, True, NAVY)
line(s, 4.78, 2.70, 4.78, 3.10, NAVY, 1.5)
line(s, 4.78, 3.86, 2.56, 4.48, SCHOOL_BLUE, 1.4)
line(s, 4.78, 3.86, 7.02, 4.48, SCHOOL_BLUE, 1.4)
text(s, "$s1 : EntitySet", 1.88, 5.48, 1.34, 0.30, 9.2, True, AUX_TEAL, PP_ALIGN.CENTER)
text(s, "$s2 : EntitySet", 6.34, 5.48, 1.34, 0.30, 9.2, True, AUX_TEAL, PP_ALIGN.CENTER)

section_label(s, "执行协议", 0.56, 6.04, 1.40)
text(s, "Schema 校验  ·  输入输出类型  ·  中间变量  ·  数据依赖  ·  错误状态",
     1.88, 5.99, 7.13, 0.38, 10.5, False, NAVY, PP_ALIGN.CENTER)
text(s, "算子集合：constraint / path / intersect / union / difference / count / enumerate / aggregate / compare / contains",
     0.76, 6.55, 8.55, 0.25, 8.5, False, GRAY, PP_ALIGN.CENTER)


# 8 Method two ---------------------------------------------------------------
s = new_slide("多粒度可执行反馈为长程检索提供局部信用信号", "方法二  ·  Credit assignment", 8)
box(s, "R = R_answer + λp R_plan + λg R_progress + λf R_format − λs C_step − λr C_repeat",
    0.72, 1.45, 8.62, 0.68, NAVY, None, 15, True, WHITE)
terms = [("Answer", "终局 EM / F1\n集合与计数正确性"),
         ("Plan", "计划可解析\n类型与依赖合法"),
         ("Progress", "支持事实覆盖\n状态向目标收敛"),
         ("Format", "JSON / 工具\n字段协议正确"),
         ("Cost", "步数 · 重复\ntoken · 延迟")]
for i, (a, b) in enumerate(terms):
    x = 0.54 + i * 1.83
    top_fill = GOLD if a == "Progress" else SCHOOL_BLUE
    body_fill = PALE_GOLD if a == "Progress" else WHITE
    border = GOLD if a == "Progress" else LINE
    box(s, a, x, 2.58, 1.56, 0.42, top_fill, None, 10.5, True, WHITE)
    box(s, b, x, 3.02, 1.56, 0.98, body_fill, border, 9.6, False, INK)

text(s, "逐步反馈", 0.82, 4.63, 1.00, 0.30, 10.5, True, AUX_TEAL, PP_ALIGN.CENTER)
arrow(s, 1.92, 4.69, 1.10, 0.20, AUX_TEAL)
box(s, "中间动作获得局部信用", 3.16, 4.50, 2.20, 0.56,
    PALE_BLUE, SCHOOL_BLUE, 10.5, True, NAVY)
arrow(s, 5.52, 4.69, 1.10, 0.20, GOLD)
box(s, "整条轨迹优化答案与成本", 6.76, 4.50, 2.30, 0.56,
    PALE_GOLD, GOLD, 10.5, True, NAVY)

shape(s, 0.78, 5.54, 8.45, 0.85, WHITE, LINE)
text(s, "反馈构造原则", 1.00, 5.78, 1.20, 0.30, 10.5, True, SCHOOL_BLUE)
text(s, "终局答案约束轨迹正确性  ·  执行等价性支持多种有效计划  ·  重复与成本项控制检索预算",
     2.28, 5.67, 6.66, 0.46, 10.3, False, GRAY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, "核心研究在于多粒度可执行反馈的构造与策略利用。",
     0.80, 6.61, 8.40, 0.30, 11, True, NAVY, PP_ALIGN.CENTER)


# 9 Training ------------------------------------------------------------------
s = new_slide("课程式训练逐步形成基础调用、组合决策与错误恢复能力", "训练流程", 9)
text(s, "现有环境与成功轨迹为训练提供冷启动基础", 0.65, 1.38, 4.30, 0.32,
     10.5, False, GRAY)
stages = [("01", "SFT 冷启动", "4,277 条成功轨迹\n学习动作协议与基本调用"),
          ("02", "基础任务优化", "单跳 · 简单路径 · 计数\n答案与格式反馈"),
          ("03", "组合任务优化", "集合 · 多约束 · 比较\n计划与进展反馈"),
          ("04", "错误恢复训练", "空结果 · 别名 · 非法关系\n修正与停止行为")]
for i, (num, a, b) in enumerate(stages):
    x = 0.48 + i * 2.35
    box(s, num, x + 0.67, 1.88, 0.54, 0.54, SCHOOL_BLUE, None, 11, True, WHITE,
        kind=MSO_SHAPE.OVAL)
    line(s, x + 0.94, 2.42, x + 0.94, 2.76, SCHOOL_BLUE, 1.1)
    text(s, a, x, 2.86, 1.88, 0.34, 11.5, True, NAVY, PP_ALIGN.CENTER)
    shape(s, x, 3.30, 1.88, 1.02, WHITE, LINE)
    text(s, b, x + 0.10, 3.48, 1.68, 0.62, 9.7, False, GRAY,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if i < 3: arrow(s, x + 1.94, 3.65, 0.28, 0.20, BLUE_GRAY)

section_label(s, "训练资源与实现", 0.54, 4.82, 2.2)
table(s, ["组成", "当前基础", "训练阶段使用方式"],
      [["知识环境", "RadarKG + 文本库 + 工具", "在线执行与状态反馈"],
       ["轨迹数据", "4,277 条成功轨迹", "SFT 冷启动"],
       ["问题数据", "Train 4,277 / Dev 582 / Test 542", "课程训练与独立评测"],
       ["优化实现", "SFT / 策略优化脚本原型", "开发集确定优化器与超参数"]],
      0.64, 5.23, 8.72, 1.34, [1.35, 3.20, 4.17], 8.8)


# 10 Existing foundation ------------------------------------------------------
s = new_slide("知识环境、复杂问答数据与执行器构成完整研究基础", "已有研究基础", 10)
for i, (v, lab) in enumerate([("22,241", "RadarKG 边"), ("10,744", "实体"),
                              ("1,432", "叙述文本段"), ("5,401", "Oracle 回放")]):
    metric(s, v, lab, 0.56 + i * 2.22, 1.42, 1.92, SCHOOL_BLUE if i < 2 else AUX_TEAL)

section_label(s, "图文检索环境", 0.54, 2.76, 2.2)
pipe = ["BM25 + FAISS", "RRF 融合", "图扩展", "Cross-Encoder 重排"]
for i, value in enumerate(pipe):
    x = 0.68 + i * 1.58
    box(s, value, x, 3.20, 1.32, 0.55, PALE_BLUE, SCHOOL_BLUE, 9.4, True, NAVY)
    if i < 3: arrow(s, x + 1.36, 3.38, 0.18, 0.17)

section_label(s, "数据与执行器", 6.95, 2.76, 2.0)
text(s, "Train / Dev / Test", 7.02, 3.20, 1.30, 0.28, 9.5, True, GRAY)
text(s, "4,277 / 582 / 542", 8.18, 3.16, 1.10, 0.34, 11.2, True, NAVY, PP_ALIGN.RIGHT)
text(s, "成功轨迹", 7.02, 3.66, 1.30, 0.28, 9.5, True, GRAY)
text(s, "4,277", 8.18, 3.62, 1.10, 0.34, 11.2, True, NAVY, PP_ALIGN.RIGHT)
text(s, "类型化算子", 7.02, 4.12, 1.30, 0.28, 9.5, True, GRAY)
text(s, "10 类", 8.18, 4.08, 1.10, 0.34, 11.2, True, NAVY, PP_ALIGN.RIGHT)

section_label(s, "原型可行性", 0.54, 4.50, 2.0)
metric(s, "100%", "计划有效率", 0.70, 4.92, 1.58, SCHOOL_BLUE)
metric(s, "100%", "执行成功率", 2.48, 4.92, 1.58, AUX_TEAL)
metric(s, "98.6%", "答案正确率（71/72）", 4.26, 4.92, 1.86, GOLD)
text(s, "小规模组合原型已验证执行链路、类型协议与工具协同的可行性。",
     6.48, 4.96, 2.70, 0.76, 10.5, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

shape(s, 0.72, 6.20, 8.52, 0.53, PALE_GOLD, GOLD)
text(s, "前期观察：固定策略在 RadarKG-QA-499 上有效，公开集收益呈现任务依赖性，进一步指向可学习的自适应检索策略。",
     0.88, 6.27, 8.20, 0.30, 9.8, True, NAVY, PP_ALIGN.CENTER)


# 11 Experiment ---------------------------------------------------------------
s = new_slide("2×2 对照独立检验动作空间设计与执行反馈学习", "实验设计", 11)
text(s, "动作空间", 0.62, 1.40, 1.05, 0.28, 9.5, True, GRAY, PP_ALIGN.CENTER)
text(s, "扁平工具调用", 3.14, 1.40, 1.80, 0.30, 11.2, True, BLUE_GRAY, PP_ALIGN.CENTER)
text(s, "类型化 / 分层决策", 5.55, 1.40, 2.02, 0.30, 11.2, True, SCHOOL_BLUE, PP_ALIGN.CENTER)
text(s, "训练方式", 1.18, 2.25, 0.85, 0.28, 9.5, True, GRAY, PP_ALIGN.CENTER)

cells = [("扁平 SFT", "成功轨迹模仿基线", 2.85, 1.88, PALE_BLUE, BLUE_GRAY),
         ("类型化 SFT", "结构化动作 + 成功轨迹", 5.42, 1.88, WHITE, SCHOOL_BLUE),
         ("扁平 RL", "扁平空间策略优化", 2.85, 3.16, PALE_BLUE, BLUE_GRAY),
         ("完整方法 RL", "类型化动作 + 多粒度反馈", 5.42, 3.16, PALE_GOLD, GOLD)]
for a, b, x, y, fill, border in cells:
    shape(s, x, y, 2.18, 0.98, fill, border, width=1.1)
    text(s, a, x + 0.08, y + 0.14, 2.02, 0.30, 11.2, True, NAVY, PP_ALIGN.CENTER)
    text(s, b, x + 0.08, y + 0.54, 2.02, 0.24, 8.9, False, GRAY, PP_ALIGN.CENTER)
text(s, "SFT", 1.65, 2.19, 0.70, 0.28, 10.5, True, GRAY, PP_ALIGN.CENTER)
text(s, "SFT + RL", 1.48, 3.49, 1.02, 0.28, 10.5, True, GRAY, PP_ALIGN.CENTER)

# Hypothesis arrows.
arrow(s, 5.07, 2.24, 0.28, 0.20, SCHOOL_BLUE)
text(s, "H1  动作空间设计", 3.88, 2.80, 1.55, 0.24, 9.2, True, SCHOOL_BLUE, PP_ALIGN.CENTER)
line(s, 6.50, 2.91, 6.50, 3.13, GOLD, 1.4)
arrow(s, 6.38, 2.99, 0.24, 0.26, GOLD, "down")
text(s, "H2  执行反馈学习", 6.70, 2.83, 1.58, 0.24, 9.2, True, GOLD)

section_label(s, "验证层", 0.54, 4.47, 1.4)
tests = [("IID", "总体效果"), ("实体隔离", "实体泛化"),
         ("组合隔离", "策略泛化", True), ("错误恢复", "利用环境观察")]
for i, item in enumerate(tests):
    a, b, *mark = item; x = 0.62 + i * 1.70
    fill = PALE_GOLD if mark else WHITE; border = GOLD if mark else LINE
    box(s, a + "\n" + b, x, 4.91, 1.48, 0.66, fill, border, 9.5, bool(mark),
        GOLD if mark else NAVY)
text(s, "公开验证", 7.52, 4.49, 1.05, 0.28, 10.5, True, AUX_TEAL, PP_ALIGN.CENTER)
box(s, "KQA Pro\n类型化程序与集合运算", 7.15, 4.91, 1.12, 0.66,
    PALE_BLUE, SCHOOL_BLUE, 8.5, True, NAVY)
box(s, "MetaQA\n1 / 2 / 3-hop 路径", 8.39, 4.91, 1.02, 0.66,
    PALE_BLUE, SCHOOL_BLUE, 8.5, True, NAVY)

shape(s, 0.76, 5.93, 8.50, 0.58, WHITE, LINE)
text(s, "分层报告：环境 oracle  →  策略执行  →  最终答案",
     1.02, 6.02, 8.00, 0.30, 11, True, NAVY, PP_ALIGN.CENTER)
text(s, "组合隔离检验模型能否将已学习的基础操作迁移到未见程序结构。",
     0.78, 6.64, 8.46, 0.28, 9.7, False, GRAY, PP_ALIGN.CENTER)


# 12 Innovations --------------------------------------------------------------
s = new_slide("两项方法创新与两个核心问题形成一一对应关系", "拟创新点", 12)
headers = ["核心问题", "方法创新", "关键机制", "主要验证"]
rows = [
    ["动作空间复杂", "类型化图检索决策建模", "宏动作 · 类型参数 · 变量依赖 · 确定性执行",
     "扁平 SFT → 类型化 SFT\n扁平 RL → 完整方法 RL"],
    ["长程信用分配", "多粒度可执行反馈策略学习", "答案 · 计划 · 进展 · 恢复 · 成本",
     "类型化 SFT → 完整方法 RL\n答案奖励 → 多粒度反馈"],
]
table(s, headers, rows, 0.58, 1.55, 8.86, 2.30, [1.48, 2.15, 2.95, 2.28], 9.2,
      SCHOOL_BLUE)

text(s, "验证关系", 0.66, 4.35, 1.05, 0.30, 10.5, True, AUX_TEAL)
flow = [("问题 1", "动作空间", BLUE_GRAY), ("创新 1", "类型化决策", SCHOOL_BLUE),
        ("证据", "2×2 横向对照", GOLD), ("问题 2", "稀疏反馈", BLUE_GRAY),
        ("创新 2", "执行反馈学习", SCHOOL_BLUE), ("证据", "2×2 纵向对照", GOLD)]
for i, (a, b, color) in enumerate(flow):
    x = 0.68 + i * 1.43
    box(s, a + "\n" + b, x, 4.82, 1.18, 0.67,
        PALE_GOLD if color == GOLD else PALE_BLUE, color, 9.2, True, NAVY)
    if i in (0, 1, 3, 4): arrow(s, x + 1.22, 5.05, 0.15, 0.16, BLUE_GRAY)

shape(s, 1.06, 5.97, 7.84, 0.67, PALE_GOLD, GOLD)
text(s, "组合隔离、错误恢复与公开数据集共同检验策略的泛化能力和适用范围",
     1.28, 6.11, 7.40, 0.34, 12, True, NAVY, PP_ALIGN.CENTER)


# 13 Plan and summary ---------------------------------------------------------
s = new_slide("研究计划围绕“建模—学习—验证—写作”形成闭环", "研究计划与总结", 13)
section_label(s, "研究进度", 0.50, 1.25, 1.6)
months = range(1, 11); left = 2.65; cw = 0.55
for j, m in enumerate(months):
    box(s, str(m), left + j * cw, 1.57, cw - 0.02, 0.34, NAVY, None, 8.3, True, WHITE)
phases = [("问题与数据冻结", 1, 2, SCHOOL_BLUE),
          ("类型化环境与 2×2 基线", 3, 4, AUX_TEAL),
          ("SFT / 策略优化 / 消融", 4, 6, GOLD),
          ("雷达与公开集实验", 7, 8, SCHOOL_BLUE),
          ("系统整合与论文写作", 9, 10, AUX_TEAL)]
for i, (name, start, end, color) in enumerate(phases):
    y = 2.08 + i * 0.47
    text(s, name, 0.62, y + 0.02, 1.82, 0.27, 9.2, False, INK)
    for m in months:
        shape(s, left + (m - 1) * cw, y, cw - 0.02, 0.30,
              color if start <= m <= end else PALE_BLUE, None)

section_label(s, "预期成果", 0.50, 4.68, 1.6, AUX_TEAL)
outs = [("方法 1", "类型化图检索决策"), ("方法 2", "多粒度反馈策略学习"),
        ("数据", "组合隔离雷达问答集"), ("实验", "RadarKG / KQA Pro / MetaQA"),
        ("系统", "可执行轨迹展示原型")]
for i, (a, b) in enumerate(outs):
    x = 0.58 + i * 1.80
    text(s, a, x, 5.14, 0.58, 0.28, 9.2, True, AUX_TEAL, PP_ALIGN.CENTER)
    text(s, b, x + 0.60, 5.08, 1.05, 0.42, 8.9, True, NAVY,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

shape(s, 0.72, 5.82, 8.52, 0.92, WHITE, SCHOOL_BLUE, width=1.1)
text(s, "为什么", 0.96, 6.05, 0.65, 0.28, 9.8, True, SCHOOL_BLUE, PP_ALIGN.CENTER)
text(s, "复杂问答需要策略学习", 1.58, 6.00, 1.88, 0.38, 10.5, True, NAVY, PP_ALIGN.CENTER)
text(s, "如何做", 3.55, 6.05, 0.65, 0.28, 9.8, True, SCHOOL_BLUE, PP_ALIGN.CENTER)
text(s, "类型化决策 + 可执行反馈", 4.15, 6.00, 2.22, 0.38, 10.5, True, NAVY, PP_ALIGN.CENTER)
text(s, "如何证", 6.47, 6.05, 0.65, 0.28, 9.8, True, SCHOOL_BLUE, PP_ALIGN.CENTER)
text(s, "2×2 对照 + 组合隔离", 7.10, 6.00, 1.82, 0.38, 10.5, True, NAVY, PP_ALIGN.CENTER)


# 14 Backup: prior observations ----------------------------------------------
s = new_slide("前期观察：固定检索策略呈现明显任务依赖性", "备份  ·  已有实验", 14, True)
section_label(s, "RadarKG-QA-499", 0.52, 1.35, 2.5)
radar = [("单轮基线", 52.7, BLUE_GRAY), ("RoG-style", 60.3, AUX_TEAL),
         ("固定策略路由", 88.6, SCHOOL_BLUE)]
for i, (name, value, color) in enumerate(radar):
    y = 1.88 + i * 0.70
    text(s, name, 0.68, y + 0.07, 1.20, 0.26, 9.8, False, INK)
    shape(s, 1.92, y, 3.05, 0.38, PALE_BLUE, None)
    shape(s, 1.92, y, 3.05 * value / 100, 0.38, color, None)
    text(s, f"{value:.1f}%", 5.05, y + 0.01, 0.58, 0.30, 10, True, color)
text(s, "固定结构化检索在领域组合题上表现出明显优势", 0.72, 4.16, 4.65, 0.34,
     10.5, True, SCHOOL_BLUE, PP_ALIGN.CENTER)

section_label(s, "KQA Pro-502", 5.37, 1.35, 2.0, AUX_TEAL)
kqa = [("Baseline", 28.3, BLUE_GRAY), ("Strategy-oracle", 28.1, AUX_TEAL)]
for i, (name, value, color) in enumerate(kqa):
    y = 2.08 + i * 0.92
    text(s, name, 5.52, y + 0.07, 1.38, 0.26, 9.8, False, INK)
    shape(s, 6.94, y, 1.78, 0.38, PALE_BLUE, None)
    shape(s, 6.94, y, 1.78 * value / 40, 0.38, color, None)
    text(s, f"{value:.1f}%", 8.80, y + 0.01, 0.58, 0.30, 10, True, color)
text(s, "Δ = −0.2 pp  [−2.4, +2.0]", 5.58, 4.16, 3.52, 0.34,
     10.5, True, AUX_TEAL, PP_ALIGN.CENTER)

shape(s, 0.78, 5.10, 8.43, 1.10, PALE_GOLD, GOLD)
text(s, "研究动机", 1.02, 5.35, 1.05, 0.28, 10.5, True, GOLD)
text(s, "领域内有效  +  跨域收益不稳定  →  从固定路由走向可学习的自适应图检索策略",
     2.00, 5.27, 6.91, 0.42, 13, True, NAVY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, "数据来源：qa_500 full 3-way；KQA Pro 502-Q e2e + paired bootstrap。",
     0.80, 6.52, 8.40, 0.25, 8.5, False, GRAY, PP_ALIGN.CENTER)


# 15 Backup: evaluation and risk ---------------------------------------------
s = new_slide("评测口径与实施风险用于保障结论可解释、可复现", "备份  ·  研究实施", 15, True)
section_label(s, "分层评测口径", 0.52, 1.32, 2.0)
table(s, ["层次", "回答的问题", "主要指标"],
      [["环境 oracle", "知识和工具能否支持金标答案", "覆盖上限 / oracle 执行率"],
       ["策略执行", "动作、参数与计划能否正确完成", "计划有效率 / 执行成功率"],
       ["最终答案", "完整系统能否返回正确答案", "EM / 集合 F1 / 计数误差"]],
      0.62, 1.78, 4.18, 2.18, [1.05, 1.95, 1.18], 8.7)

section_label(s, "实施风险与应对", 5.12, 1.32, 2.2, GOLD)
table(s, ["风险", "应对"],
      [["策略优化增益有限", "强化组合与恢复任务；分析反馈方差"],
       ["过程奖励投机", "答案门控、重复惩罚与轨迹审计"],
       ["KG 覆盖产生噪声", "闭世界训练子集；报告 oracle 上限"],
       ["计算资源约束", "小模型 LoRA、量化与 offload"]],
      5.18, 1.78, 4.18, 2.18, [1.45, 2.73], 8.2, GOLD)

section_label(s, "统计与复现", 0.52, 4.48, 1.8, SCHOOL_BLUE)
bullets(s, ["核心实验至少运行 3 个随机种子，并使用题级配对 bootstrap 或随机化检验",
            "统一基座模型、知识环境、最大步数、检索预算和答案生成器",
            "按题型、计划深度、环境覆盖和错误类别报告结果"],
        0.70, 4.92, 4.10, 1.25, 10.5)

section_label(s, "研究范围", 5.12, 4.48, 1.8, AUX_TEAL)
bullets(s, ["论文聚焦复杂问答的图检索决策与策略学习",
            "来源和置信度用于可追溯展示及扩展实验",
            "系统实现用于方法验证与应用案例展示"],
        5.30, 4.92, 3.88, 1.25, 10.5)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUTPUT))
print(f"saved: {OUTPUT}")
print(f"slides: {len(prs.slides)}; main=13; backup=2; size={SW/914400:.2f}x{SH/914400:.2f}")
